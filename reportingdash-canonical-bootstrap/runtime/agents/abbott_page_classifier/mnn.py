"""Deterministic, offline Abbott MNN workbook parsing.

The source calls the field ``МНН`` but currently contains both international
names and approved product labels.  This module preserves those labels as
source evidence and performs only structural normalization; it never applies
medical or fuzzy corrections.
"""

from __future__ import annotations

from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook

from .normalization import normalize_title, normalize_url, sha256_text


_SOURCE_ID_RE = re.compile(r"\s*\[\d+\]\s*")
_DOSAGE_SLASH_RE = re.compile(r"(?P<left>\d+)\s*/\s*(?P<right>\d+)")
_NA_SLASH_RE = re.compile(r"\bn\s*/\s*a\b", re.IGNORECASE)
_LIST_SEPARATOR_RE = re.compile(r"[,\n\r/]+")
_DOSAGE_SLASH_SENTINEL = "\u0000dosage-slash\u0000"
_NA_SLASH_SENTINEL = "\u0000na-slash\u0000"
_PLACEHOLDER_KEYS = frozenset(
    {
        "-",
        "—",
        "#n/a",
        "n/a",
        "na",
        "не найдено",
        "нет",
        "нет бренда",
        "нет доступа",
        "нет материала",
        "ошибка",
        "подраздел",
        "раздел",
        "скрыто",
    }
)


def _canonical_json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _header_key(value: object) -> str:
    return "".join(character for character in normalize_title(str(value or "")).casefold() if character.isalnum())


@dataclass(frozen=True)
class NormalizedMnn:
    key: str
    label: str


@dataclass(frozen=True)
class MnnProvenance:
    source_sheet: str
    source_row_ordinal: int
    raw_url: str
    raw_value: str
    source_row_fingerprint: str


@dataclass(frozen=True)
class MnnClaim:
    normalized_url: str
    mnn_key: str
    mnn_label: str
    provenance: tuple[MnnProvenance, ...]


@dataclass(frozen=True)
class UnlinkedMnnRow:
    source_sheet: str
    source_row_ordinal: int
    raw_value: str
    mnn_labels: tuple[str, ...]
    source_row_fingerprint: str


@dataclass(frozen=True)
class RejectedMnnValue:
    source_sheet: str
    source_row_ordinal: int
    raw_value: str
    rejected_label: str
    reason_code: str
    source_row_fingerprint: str


@dataclass(frozen=True)
class MnnSnapshot:
    source_sha256: str
    mnn_row_count: int
    linked_row_count: int
    unlinked_row_count: int
    rejected_value_count: int
    claims: tuple[MnnClaim, ...]
    unlinked_rows: tuple[UnlinkedMnnRow, ...]
    rejected_values: tuple[RejectedMnnValue, ...]


@dataclass(frozen=True)
class MnnEntityMapping:
    content_entity_id: int
    normalized_url: str
    mnn_key: str
    mnn_label: str
    claim_fingerprint: str
    mapping_fingerprint: str


@dataclass(frozen=True)
class MnnResolutionReport:
    mappings: tuple[MnnEntityMapping, ...]
    unresolved: tuple[MnnClaim, ...]
    collisions: tuple[MnnClaim, ...]
    unlinked_count: int


def normalize_mnn_label(value: object) -> NormalizedMnn:
    """Normalize one already-separated source label without guessing aliases."""

    label = normalize_title(_SOURCE_ID_RE.sub(" ", str(value or ""))).strip(" ,/")
    key = normalize_title(label.replace("®", "").replace("™", "")).casefold()
    return NormalizedMnn(key=key, label=label)


def _split_mnn_values(
    value: object,
) -> tuple[tuple[str, ...], tuple[tuple[str, str], ...]]:
    """Split source values and separate explicit non-MNN placeholders."""

    raw = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    if not raw.strip():
        return (), ()
    protected = _DOSAGE_SLASH_RE.sub(
        lambda match: f"{match.group('left')}{_DOSAGE_SLASH_SENTINEL}{match.group('right')}",
        raw,
    )
    protected = _NA_SLASH_RE.sub(_NA_SLASH_SENTINEL, protected)
    labels: list[str] = []
    rejected: list[tuple[str, str]] = []
    seen: set[str] = set()
    for part in _LIST_SEPARATOR_RE.split(protected):
        normalized = normalize_mnn_label(
            part.replace(_DOSAGE_SLASH_SENTINEL, "/").replace(_NA_SLASH_SENTINEL, "n/a")
        )
        if not normalized.key or normalized.key in seen:
            continue
        seen.add(normalized.key)
        if normalized.key in _PLACEHOLDER_KEYS:
            rejected.append((normalized.label, "MNN_PLACEHOLDER"))
            continue
        if (
            len(normalized.label) > 500
            or len(normalized.key) > 255
            or not _header_key(normalized.label)
        ):
            rejected.append((normalized.label, "MNN_INVALID"))
            continue
        labels.append(normalized.label)
    return tuple(labels), tuple(rejected)


def split_mnn_labels(value: object) -> tuple[str, ...]:
    """Split approved product lists while keeping numeric dosage slashes."""

    return _split_mnn_values(value)[0]


def _sheet_header(sheet: object) -> tuple[int, int, int] | None:
    for row_ordinal, values in enumerate(sheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        positions = {_header_key(value): index for index, value in enumerate(values)}
        mnn_col = positions.get("мнн")
        url_col = next((positions[key] for key in ("ссылка", "url") if key in positions), None)
        if mnn_col is not None and url_col is not None:
            return row_ordinal, url_col, mnn_col
    return None


def _row_fingerprint(*, sheet: str, row: int, raw_url: str, raw_value: str) -> str:
    return sha256_text(
        _canonical_json(
            {
                "raw_url": raw_url,
                "raw_value": raw_value,
                "source_row_ordinal": row,
                "source_sheet": sheet,
            }
        )
    )


def read_mnn_workbook(path: Path) -> MnnSnapshot:
    """Read all sheets with exact URL/MNN headers into stable multi-value claims."""

    source_path = Path(path)
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    grouped: dict[tuple[str, str], tuple[str, list[MnnProvenance]]] = {}
    unlinked: list[UnlinkedMnnRow] = []
    rejected: list[RejectedMnnValue] = []
    mnn_row_count = 0
    linked_row_count = 0
    recognized_sheets = 0
    try:
        for sheet in workbook.worksheets:
            header = _sheet_header(sheet)
            if header is None:
                continue
            recognized_sheets += 1
            header_row, url_col, mnn_col = header
            for row_ordinal, values in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                raw_value = str(values[mnn_col] or "") if mnn_col < len(values) else ""
                labels, rejected_labels = _split_mnn_values(raw_value)
                if not labels and not rejected_labels:
                    continue
                mnn_row_count += 1
                raw_url = normalize_title(str(values[url_col] or "")) if url_col < len(values) else ""
                fingerprint = _row_fingerprint(
                    sheet=sheet.title,
                    row=row_ordinal,
                    raw_url=raw_url,
                    raw_value=raw_value,
                )
                for rejected_label, reason_code in rejected_labels:
                    rejected.append(
                        RejectedMnnValue(
                            source_sheet=sheet.title,
                            source_row_ordinal=row_ordinal,
                            raw_value=raw_value,
                            rejected_label=rejected_label,
                            reason_code=reason_code,
                            source_row_fingerprint=fingerprint,
                        )
                    )
                if not labels:
                    continue
                normalized_url = normalize_url(raw_url).value
                if not normalized_url:
                    unlinked.append(
                        UnlinkedMnnRow(
                            source_sheet=sheet.title,
                            source_row_ordinal=row_ordinal,
                            raw_value=raw_value,
                            mnn_labels=labels,
                            source_row_fingerprint=fingerprint,
                        )
                    )
                    continue
                linked_row_count += 1
                provenance = MnnProvenance(
                    source_sheet=sheet.title,
                    source_row_ordinal=row_ordinal,
                    raw_url=raw_url,
                    raw_value=raw_value,
                    source_row_fingerprint=fingerprint,
                )
                for label in labels:
                    normalized = normalize_mnn_label(label)
                    pair = (normalized_url, normalized.key)
                    if pair not in grouped:
                        grouped[pair] = (normalized.label, [])
                    grouped[pair][1].append(provenance)
    finally:
        workbook.close()

    if recognized_sheets == 0:
        raise ValueError("MNN_SOURCE_HEADERS_MISSING")

    claims = tuple(
        MnnClaim(
            normalized_url=url,
            mnn_key=key,
            mnn_label=label,
            provenance=tuple(
                sorted(values, key=lambda item: (item.source_sheet, item.source_row_ordinal, item.source_row_fingerprint))
            ),
        )
        for (url, key), (label, values) in sorted(grouped.items())
    )
    return MnnSnapshot(
        source_sha256=hashlib.sha256(source_path.read_bytes()).hexdigest(),
        mnn_row_count=mnn_row_count,
        linked_row_count=linked_row_count,
        unlinked_row_count=len(unlinked),
        rejected_value_count=len(rejected),
        claims=claims,
        unlinked_rows=tuple(sorted(unlinked, key=lambda item: (item.source_sheet, item.source_row_ordinal))),
        rejected_values=tuple(
            sorted(rejected, key=lambda item: (item.source_sheet, item.source_row_ordinal, item.rejected_label))
        ),
    )


def _claim_fingerprint(claim: MnnClaim, source_sha256: str) -> str:
    return sha256_text(
        _canonical_json(
            {
                "mnn_key": claim.mnn_key,
                "normalized_url": claim.normalized_url,
                "source_sha256": source_sha256,
                "source_rows": [item.source_row_fingerprint for item in claim.provenance],
            }
        )
    )


def resolve_mnn_claims(
    snapshot: MnnSnapshot,
    strong_url_owners: dict[str, tuple[int, ...]],
) -> MnnResolutionReport:
    """Resolve claims only when exact strong URL evidence has one owner."""

    owner_sets: dict[str, set[int]] = {}
    for raw_url, raw_owners in strong_url_owners.items():
        normalized_url = normalize_url(raw_url).value
        if not normalized_url:
            continue
        owners = owner_sets.setdefault(normalized_url, set())
        for raw_owner in raw_owners:
            owner = int(raw_owner)
            if owner <= 0:
                raise ValueError("MNN_STRONG_OWNER_INVALID")
            owners.add(owner)

    mappings: list[MnnEntityMapping] = []
    unresolved: list[MnnClaim] = []
    collisions: list[MnnClaim] = []
    for claim in snapshot.claims:
        owners = owner_sets.get(claim.normalized_url, set())
        if not owners:
            unresolved.append(claim)
            continue
        if len(owners) != 1:
            collisions.append(claim)
            continue
        entity_id = next(iter(owners))
        claim_fingerprint = _claim_fingerprint(claim, snapshot.source_sha256)
        mappings.append(
            MnnEntityMapping(
                content_entity_id=entity_id,
                normalized_url=claim.normalized_url,
                mnn_key=claim.mnn_key,
                mnn_label=claim.mnn_label,
                claim_fingerprint=claim_fingerprint,
                mapping_fingerprint=sha256_text(
                    _canonical_json(
                        {
                            "claim_fingerprint": claim_fingerprint,
                            "content_entity_id": entity_id,
                            "mnn_key": claim.mnn_key,
                        }
                    )
                ),
            )
        )
    return MnnResolutionReport(
        mappings=tuple(sorted(mappings, key=lambda item: (item.content_entity_id, item.mnn_key))),
        unresolved=tuple(unresolved),
        collisions=tuple(collisions),
        unlinked_count=len(snapshot.unlinked_rows),
    )
