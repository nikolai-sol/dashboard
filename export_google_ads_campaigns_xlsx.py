#!/usr/bin/env python3
"""Export detailed Google Ads campaign data to an XLSX workbook."""

from __future__ import annotations

import argparse
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from google_ads_api_client import env_first, google_ads_client, normalize_customer_id

load_dotenv(Path(__file__).parent / ".env")


DEFAULT_CAMPAIGN_IDS = ["23822556855", "23831678560"]  # DE, AT1


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--customer-id", default=env_first("GOOGLE_ADS_CUSTOMER_IDS").split(",")[0])
    parser.add_argument("--campaign-ids", default=",".join(DEFAULT_CAMPAIGN_IDS))
    parser.add_argument("--date-from", default="2024-01-01")
    parser.add_argument("--date-to", default=date.today().strftime("%Y-%m-%d"))
    parser.add_argument("--output", default="")
    return parser.parse_args()


def micros_to_units(value: Any) -> float:
    try:
        return round(float(value or 0) / 1_000_000, 6)
    except (TypeError, ValueError):
        return 0.0


def enum_name(value: Any) -> str:
    name = getattr(value, "name", None)
    return str(name if name is not None else value or "")


def scalar(value: Any) -> Any:
    if value is None:
        return ""
    if hasattr(value, "name"):
        return enum_name(value)
    return value


def campaign_filter(campaign_ids: list[str]) -> str:
    ids = ", ".join(campaign_ids)
    return f"campaign.id IN ({ids})"


def single_campaign_filter(campaign_id: str) -> str:
    return f"campaign.id = {campaign_id}"


def month_chunks(date_from: str, date_to: str) -> list[tuple[str, str]]:
    start = datetime.strptime(date_from, "%Y-%m-%d").date()
    end = datetime.strptime(date_to, "%Y-%m-%d").date()
    chunks: list[tuple[str, str]] = []
    cursor = start
    while cursor <= end:
        if cursor.month == 12:
            next_month = date(cursor.year + 1, 1, 1)
        else:
            next_month = date(cursor.year, cursor.month + 1, 1)
        chunk_end = min(end, next_month - timedelta(days=1))
        chunks.append((cursor.strftime("%Y-%m-%d"), chunk_end.strftime("%Y-%m-%d")))
        cursor = chunk_end + timedelta(days=1)
    return chunks


def add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for column, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row_index, row in enumerate(rows, 2):
        for column, header in enumerate(headers, 1):
            ws.cell(row=row_index, column=column, value=row.get(header, ""))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column, header in enumerate(headers, 1):
        max_len = min(
            max([len(str(header))] + [len(str(row.get(header, ""))) for row in rows[:500]]) + 2,
            60,
        )
        ws.column_dimensions[get_column_letter(column)].width = max_len


def query_rows(client, customer_id: str, query: str) -> list[Any]:
    service = client.get_service("GoogleAdsService")
    rows: list[Any] = []
    for batch in service.search_stream(customer_id=normalize_customer_id(customer_id), query=query):
        rows.extend(batch.results)
    return rows


def query_chunked(
    client,
    customer_id: str,
    label: str,
    campaign_ids: list[str],
    date_from: str,
    date_to: str,
    build_query,
) -> list[Any]:
    rows: list[Any] = []
    chunks = month_chunks(date_from, date_to)
    for campaign_id in campaign_ids:
        for chunk_from, chunk_to in chunks:
            query = build_query(single_campaign_filter(campaign_id), chunk_from, chunk_to)
            chunk_rows = query_rows(client, customer_id, query)
            rows.extend(chunk_rows)
        print(f"  {label}: campaign {campaign_id} done ({len(rows)} cumulative rows)")
    return rows


def main() -> None:
    args = parse_args()
    customer_id = normalize_customer_id(args.customer_id)
    campaign_ids = ["".join(ch for ch in item if ch.isdigit()) for item in args.campaign_ids.split(",")]
    campaign_ids = [item for item in campaign_ids if item]
    if not campaign_ids:
        raise SystemExit("No campaign ids provided")

    date_from = args.date_from
    date_to = args.date_to
    output = args.output or f"google_ads_DE_AT1_{date_from}_to_{date_to}.xlsx"
    output_path = Path(output).expanduser().resolve()

    client = google_ads_client()
    filt = campaign_filter(campaign_ids)

    campaign_meta_query = f"""
        SELECT
          customer.id,
          customer.descriptive_name,
          customer.currency_code,
          customer.time_zone,
          campaign.id,
          campaign.name,
          campaign.status,
          campaign.advertising_channel_type,
          campaign_budget.amount_micros
        FROM campaign
        WHERE {filt}
        ORDER BY campaign.id
    """
    def campaign_daily_query(filter_clause: str, chunk_from: str, chunk_to: str) -> str:
        return f"""
        SELECT
          segments.date,
          customer.id,
          customer.currency_code,
          campaign.id,
          campaign.name,
          campaign.status,
          campaign.advertising_channel_type,
          metrics.impressions,
          metrics.clicks,
          metrics.cost_micros,
          metrics.conversions,
          metrics.conversions_value,
          metrics.all_conversions,
          metrics.all_conversions_value
        FROM campaign
        WHERE segments.date BETWEEN '{chunk_from}' AND '{chunk_to}'
          AND {filter_clause}
        ORDER BY segments.date, campaign.id
    """

    def ad_group_daily_query(filter_clause: str, chunk_from: str, chunk_to: str) -> str:
        return f"""
        SELECT
          segments.date,
          customer.id,
          campaign.id,
          campaign.name,
          ad_group.id,
          ad_group.name,
          ad_group.status,
          ad_group.type,
          metrics.impressions,
          metrics.clicks,
          metrics.cost_micros,
          metrics.conversions,
          metrics.conversions_value,
          metrics.all_conversions,
          metrics.all_conversions_value
        FROM ad_group
        WHERE segments.date BETWEEN '{chunk_from}' AND '{chunk_to}'
          AND {filter_clause}
        ORDER BY segments.date, campaign.id, ad_group.id
    """

    def search_terms_query(filter_clause: str, chunk_from: str, chunk_to: str) -> str:
        return f"""
        SELECT
          segments.date,
          customer.id,
          customer.currency_code,
          campaign.id,
          campaign.name,
          ad_group.id,
          ad_group.name,
          search_term_view.search_term,
          search_term_view.status,
          metrics.impressions,
          metrics.clicks,
          metrics.cost_micros,
          metrics.conversions,
          metrics.conversions_value,
          metrics.all_conversions,
          metrics.all_conversions_value
        FROM search_term_view
        WHERE segments.date BETWEEN '{chunk_from}' AND '{chunk_to}'
          AND {filter_clause}
        ORDER BY segments.date, campaign.id, ad_group.id
    """

    def products_query(filter_clause: str, chunk_from: str, chunk_to: str) -> str:
        return f"""
        SELECT
          segments.date,
          customer.id,
          customer.currency_code,
          campaign.id,
          campaign.name,
          ad_group.id,
          ad_group.name,
          segments.product_item_id,
          segments.product_title,
          segments.product_brand,
          segments.product_type_l1,
          segments.product_type_l2,
          segments.product_channel,
          segments.product_merchant_id,
          metrics.impressions,
          metrics.clicks,
          metrics.cost_micros,
          metrics.conversions,
          metrics.conversions_value,
          metrics.all_conversions,
          metrics.all_conversions_value
        FROM shopping_performance_view
        WHERE segments.date BETWEEN '{chunk_from}' AND '{chunk_to}'
          AND {filter_clause}
        ORDER BY segments.date, campaign.id, ad_group.id, segments.product_item_id
    """

    def keywords_query(filter_clause: str, chunk_from: str, chunk_to: str) -> str:
        return f"""
        SELECT
          segments.date,
          customer.id,
          customer.currency_code,
          campaign.id,
          campaign.name,
          ad_group.id,
          ad_group.name,
          ad_group_criterion.criterion_id,
          ad_group_criterion.status,
          ad_group_criterion.keyword.text,
          ad_group_criterion.keyword.match_type,
          metrics.impressions,
          metrics.clicks,
          metrics.cost_micros,
          metrics.conversions,
          metrics.conversions_value,
          metrics.all_conversions,
          metrics.all_conversions_value
        FROM keyword_view
        WHERE segments.date BETWEEN '{chunk_from}' AND '{chunk_to}'
          AND {filter_clause}
        ORDER BY segments.date, campaign.id, ad_group.id, ad_group_criterion.criterion_id
    """

    print("Fetching campaign metadata...")
    meta_rows = query_rows(client, customer_id, campaign_meta_query)
    print("Fetching campaign daily metrics...")
    campaign_daily_rows = query_chunked(client, customer_id, "campaign daily", campaign_ids, date_from, date_to, campaign_daily_query)
    print("Fetching ad group daily metrics...")
    ad_group_rows = query_chunked(client, customer_id, "ad group daily", campaign_ids, date_from, date_to, ad_group_daily_query)
    print("Fetching search terms...")
    search_rows = query_chunked(client, customer_id, "search terms", campaign_ids, date_from, date_to, search_terms_query)
    print("Fetching shopping product performance...")
    product_rows = query_chunked(client, customer_id, "products", campaign_ids, date_from, date_to, products_query)
    print("Fetching keyword performance...")
    keyword_rows = query_chunked(client, customer_id, "keywords", campaign_ids, date_from, date_to, keywords_query)

    def campaign_meta(row: Any) -> dict[str, Any]:
        return {
            "customer_id": str(row.customer.id),
            "customer_name": row.customer.descriptive_name,
            "currency_code": row.customer.currency_code,
            "time_zone": row.customer.time_zone,
            "campaign_id": str(row.campaign.id),
            "campaign_name": row.campaign.name,
            "campaign_status": enum_name(row.campaign.status),
            "channel_type": enum_name(row.campaign.advertising_channel_type),
            "daily_budget": micros_to_units(row.campaign_budget.amount_micros),
        }

    def metric_fields(row: Any) -> dict[str, Any]:
        cost = micros_to_units(row.metrics.cost_micros)
        clicks = int(row.metrics.clicks or 0)
        impressions = int(row.metrics.impressions or 0)
        conversions = float(row.metrics.conversions or 0)
        conversion_value = float(row.metrics.conversions_value or 0)
        return {
            "impressions": impressions,
            "clicks": clicks,
            "cost": cost,
            "ctr_pct": round((clicks / impressions) * 100, 6) if impressions else "",
            "avg_cpc": round(cost / clicks, 6) if clicks else "",
            "conversions": conversions,
            "conversion_value": conversion_value,
            "cost_per_conversion": round(cost / conversions, 6) if conversions else "",
            "roas": round(conversion_value / cost, 6) if cost else "",
            "all_conversions": float(row.metrics.all_conversions or 0),
            "all_conversion_value": float(row.metrics.all_conversions_value or 0),
        }

    campaign_daily = [
        {
            "date": str(row.segments.date),
            "customer_id": str(row.customer.id),
            "currency_code": row.customer.currency_code,
            "campaign_id": str(row.campaign.id),
            "campaign_name": row.campaign.name,
            "campaign_status": enum_name(row.campaign.status),
            "channel_type": enum_name(row.campaign.advertising_channel_type),
            **metric_fields(row),
        }
        for row in campaign_daily_rows
    ]

    ad_group_daily = [
        {
            "date": str(row.segments.date),
            "customer_id": str(row.customer.id),
            "campaign_id": str(row.campaign.id),
            "campaign_name": row.campaign.name,
            "ad_group_id": str(row.ad_group.id),
            "ad_group_name": row.ad_group.name,
            "ad_group_status": enum_name(row.ad_group.status),
            "ad_group_type": enum_name(row.ad_group.type),
            **metric_fields(row),
        }
        for row in ad_group_rows
    ]

    search_terms = [
        {
            "date": str(row.segments.date),
            "customer_id": str(row.customer.id),
            "currency_code": row.customer.currency_code,
            "campaign_id": str(row.campaign.id),
            "campaign_name": row.campaign.name,
            "ad_group_id": str(row.ad_group.id),
            "ad_group_name": row.ad_group.name,
            "search_term": row.search_term_view.search_term,
            "search_term_status": enum_name(row.search_term_view.status),
            **metric_fields(row),
        }
        for row in search_rows
    ]

    products = [
        {
            "date": str(row.segments.date),
            "customer_id": str(row.customer.id),
            "currency_code": row.customer.currency_code,
            "campaign_id": str(row.campaign.id),
            "campaign_name": row.campaign.name,
            "ad_group_id": str(row.ad_group.id),
            "ad_group_name": row.ad_group.name,
            "product_item_id": row.segments.product_item_id,
            "product_title": row.segments.product_title,
            "product_brand": row.segments.product_brand,
            "product_type_l1": row.segments.product_type_l1,
            "product_type_l2": row.segments.product_type_l2,
            "product_channel": row.segments.product_channel,
            "merchant_id": str(row.segments.product_merchant_id or ""),
            **metric_fields(row),
        }
        for row in product_rows
    ]

    keywords = [
        {
            "date": str(row.segments.date),
            "customer_id": str(row.customer.id),
            "currency_code": row.customer.currency_code,
            "campaign_id": str(row.campaign.id),
            "campaign_name": row.campaign.name,
            "ad_group_id": str(row.ad_group.id),
            "ad_group_name": row.ad_group.name,
            "criterion_id": str(row.ad_group_criterion.criterion_id),
            "criterion_status": enum_name(row.ad_group_criterion.status),
            "keyword_text": row.ad_group_criterion.keyword.text,
            "match_type": enum_name(row.ad_group_criterion.keyword.match_type),
            **metric_fields(row),
        }
        for row in keyword_rows
    ]

    summary_by_campaign: dict[str, dict[str, Any]] = {}
    for row in campaign_daily:
        key = row["campaign_id"]
        target = summary_by_campaign.setdefault(
            key,
            {
                "campaign_id": row["campaign_id"],
                "campaign_name": row["campaign_name"],
                "impressions": 0,
                "clicks": 0,
                "cost": 0.0,
                "conversions": 0.0,
                "conversion_value": 0.0,
            },
        )
        target["impressions"] += int(row["impressions"] or 0)
        target["clicks"] += int(row["clicks"] or 0)
        target["cost"] += float(row["cost"] or 0)
        target["conversions"] += float(row["conversions"] or 0)
        target["conversion_value"] += float(row["conversion_value"] or 0)
    summary = []
    for row in summary_by_campaign.values():
        row["cost"] = round(row["cost"], 6)
        row["conversion_value"] = round(row["conversion_value"], 6)
        row["ctr_pct"] = round((row["clicks"] / row["impressions"]) * 100, 6) if row["impressions"] else ""
        row["avg_cpc"] = round(row["cost"] / row["clicks"], 6) if row["clicks"] else ""
        row["cost_per_conversion"] = round(row["cost"] / row["conversions"], 6) if row["conversions"] else ""
        row["roas"] = round(row["conversion_value"] / row["cost"], 6) if row["cost"] else ""
        summary.append(row)

    wb = Workbook()
    del wb[wb.sheetnames[0]]
    add_sheet(
        wb,
        "read_me",
        ["field", "value"],
        [
            {"field": "source", "value": "Google Ads API live export"},
            {"field": "customer_id", "value": customer_id},
            {"field": "campaign_ids", "value": ", ".join(campaign_ids)},
            {"field": "date_from", "value": date_from},
            {"field": "date_to", "value": date_to},
            {"field": "note", "value": "Shopping campaigns usually have search terms and product performance, not manual keyword rows."},
        ],
    )
    add_sheet(wb, "campaign_metadata", list(campaign_meta(meta_rows[0]).keys()) if meta_rows else ["empty"], [campaign_meta(row) for row in meta_rows])
    add_sheet(wb, "campaign_summary", list(summary[0].keys()) if summary else ["empty"], summary)
    add_sheet(wb, "campaign_daily", list(campaign_daily[0].keys()) if campaign_daily else ["empty"], campaign_daily)
    add_sheet(wb, "ad_group_daily", list(ad_group_daily[0].keys()) if ad_group_daily else ["empty"], ad_group_daily)
    add_sheet(wb, "search_terms_daily", list(search_terms[0].keys()) if search_terms else ["empty"], search_terms)
    add_sheet(wb, "product_performance_daily", list(products[0].keys()) if products else ["empty"], products)
    add_sheet(wb, "keyword_performance_daily", list(keywords[0].keys()) if keywords else ["empty"], keywords)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Wrote {output_path}")
    print(
        "Rows: "
        f"metadata={len(meta_rows)}, campaign_daily={len(campaign_daily)}, "
        f"ad_group_daily={len(ad_group_daily)}, search_terms={len(search_terms)}, "
        f"products={len(products)}, keywords={len(keywords)}"
    )


if __name__ == "__main__":
    main()
