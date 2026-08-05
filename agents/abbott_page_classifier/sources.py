"""Deterministic, offline readers for captured Abbott registry sources."""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
import json
from pathlib import Path
from types import MappingProxyType
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

from .domain import CanonicalClassification, MaterialCandidate
from .normalization import (
    normalize_taxonomy_label,
    normalize_title,
    normalize_url,
    sha256_text,
)


def _canonical_json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_title(str(value))


def _header_key(value: object) -> str:
    return "".join(character for character in _text(value).casefold() if character.isalnum())


_FIELD_ALIASES: Mapping[str, tuple[str, ...]] = MappingProxyType(
    {
        "material_id": ("materialid", "идматериала", "id"),
        "title": ("title", "название", "наименованиематериала"),
        "url": ("url", "ссылка", "ссылканарегистрацию", "файлpdf"),
        "direction": ("direction", "направление", "направления"),
        "access": ("access", "доступ"),
        "material_type": ("materialtype", "типконтента", "типматериала"),
        "lifecycle": ("lifecycle", "pagestatus", "статус", "активность"),
    }
)

_SHEET_DIRECTIONS: Mapping[str, str] = MappingProxyType(
    {
        "кардио": "cardiology",
        "cardio": "cardiology",
        "гастро": "gastroenterology",
        "gastro": "gastroenterology",
        "невро": "neurology_psychiatry",
        "nevro": "neurology_psychiatry",
        "респи": "respiratory_health",
        "respiratory": "respiratory_health",
        "женское": "womens_health",
        "dermatology": "dermatology",
        "дерматология": "dermatology",
        "фармацевты": "pharmacists",
    }
)


@dataclass(frozen=True)
class SourceProvenance:
    """One raw source-row occurrence, retained when candidates collapse."""

    source_name: str
    source_row_id: str
    source_fingerprint: str


@dataclass(frozen=True)
class SourceIdentityVariant:
    """Identity evidence from one occurrence in a collapsed source group."""

    source_row_id: str
    material_id: str | None
    normalized_url: str
    normalized_title: str
    material_type_code: str | None
    raw_material_type: str = ""
    raw_status: str = ""


@dataclass(frozen=True)
class _CandidateOccurrence:
    candidate: MaterialCandidate
    raw_material_type: str
    raw_status: str


@dataclass(frozen=True)
class SourceCandidate:
    """A Task 1 candidate plus its stable source key and all row provenance."""

    key: str
    candidate: MaterialCandidate
    provenance: tuple[SourceProvenance, ...]
    identity_variants: tuple[SourceIdentityVariant, ...]

    @property
    def source_name(self) -> str:
        return self.candidate.source_name

    @property
    def source_row_id(self) -> str:
        return self.candidate.source_row_id

    @property
    def title(self) -> str:
        return self.candidate.title

    @property
    def url(self) -> str:
        return self.candidate.url

    @property
    def material_id(self) -> str | None:
        return self.candidate.material_id

    @property
    def direction_code(self) -> str | None:
        return self.candidate.direction_code

    @property
    def material_type_code(self) -> str | None:
        return self.candidate.material_type_code

    @property
    def access_code(self) -> str | None:
        return self.candidate.access_code

    @property
    def lifecycle_code(self) -> str:
        return self.candidate.lifecycle_code

    @property
    def source_fingerprint(self) -> str:
        return self.candidate.source_fingerprint


@dataclass(frozen=True)
class RejectedSourceRow:
    source_name: str
    source_row_id: str
    reason_code: str
    source_fingerprint: str


@dataclass(frozen=True)
class SourceSnapshot:
    source_name: str
    source_hash: str
    source_row_count: int
    candidates: tuple[SourceCandidate, ...]
    rejected_rows: tuple[RejectedSourceRow, ...]
    duplicate_collapsed_count: int = 0
    candidates_by_key: Mapping[str, SourceCandidate] = field(init=False, repr=False)

    def __post_init__(self) -> None:
        keys = {candidate.key: candidate for candidate in self.candidates}
        if len(keys) != len(self.candidates):
            raise ValueError("DUPLICATE_CANDIDATE_KEY")
        object.__setattr__(self, "candidates_by_key", MappingProxyType(keys))
        if self.outcome_count != self.source_row_count:
            raise ValueError("SOURCE_ROW_COUNT_MISMATCH")

    @property
    def outcome_count(self) -> int:
        return len(self.candidates) + len(self.rejected_rows) + self.duplicate_collapsed_count


def _source_key(candidate: MaterialCandidate) -> str | None:
    if candidate.material_id:
        return f"material:{candidate.material_id.casefold()}"
    if candidate.url:
        return f"url:{candidate.url}"
    if candidate.title:
        material_type = candidate.material_type_code or ""
        return f"title_type:{sha256_text(normalize_title(candidate.title).casefold() + ':' + material_type)}"
    return None


def _lifecycle(raw: str) -> str:
    normalized = normalize_taxonomy_label("lifecycle", raw)
    if normalized:
        return normalized
    value = _text(raw).casefold()
    if value in {"", "да", "yes", "active", "активен"}:
        return "active"
    if value in {"нет", "no", "inactive", "неактивен", "архив"}:
        return "archive_candidate"
    return "unknown"


def _field_values(headers: Sequence[object], values: Sequence[object]) -> dict[str, str]:
    header_positions = {_header_key(header): index for index, header in enumerate(headers)}
    result: dict[str, str] = {}
    for field, aliases in _FIELD_ALIASES.items():
        position = next((header_positions[alias] for alias in aliases if alias in header_positions), None)
        result[field] = _text(values[position]) if position is not None and position < len(values) else ""
    return result


def _build_candidate(
    *,
    source_name: str,
    source_row_id: str,
    values: Mapping[str, str],
    fallback_direction: str | None = None,
) -> _CandidateOccurrence:
    title = values["title"]
    normalized_url = normalize_url(values["url"])
    material_id = values["material_id"] or None
    payload = {
        "source_name": source_name,
        "source_row_id": source_row_id,
        "title": title,
        "url": normalized_url.value,
        "material_id": material_id,
        "direction_code": normalize_taxonomy_label("direction", values["direction"])
        or fallback_direction,
        "material_type_code": normalize_taxonomy_label(
            "material_type", values["material_type"]
        ),
        "access_code": normalize_taxonomy_label("access", values["access"]),
        "lifecycle_code": _lifecycle(values["lifecycle"]),
    }
    fingerprint_payload = {
        **payload,
        "raw_material_type": values["material_type"],
        "raw_status": values["lifecycle"],
    }
    return _CandidateOccurrence(
        candidate=MaterialCandidate(
            **payload,
            source_fingerprint=sha256_text(_canonical_json(fingerprint_payload)),
        ),
        raw_material_type=values["material_type"],
        raw_status=values["lifecycle"],
    )


def _snapshot_from_candidates(
    source_name: str,
    raw_row_count: int,
    candidates: Iterable[_CandidateOccurrence],
    rejected_rows: Iterable[RejectedSourceRow],
) -> SourceSnapshot:
    grouped: dict[str, SourceCandidate] = {}
    duplicate_collapsed_count = 0
    for occurrence in candidates:
        candidate = occurrence.candidate
        key = _source_key(candidate)
        if key is None:
            raise ValueError("CANDIDATE_WITHOUT_IDENTITY")
        provenance = SourceProvenance(
            source_name=candidate.source_name,
            source_row_id=candidate.source_row_id,
            source_fingerprint=candidate.source_fingerprint,
        )
        identity_variant = SourceIdentityVariant(
            source_row_id=candidate.source_row_id,
            material_id=candidate.material_id,
            normalized_url=normalize_url(candidate.url).value,
            normalized_title=normalize_title(candidate.title),
            material_type_code=candidate.material_type_code,
            raw_material_type=occurrence.raw_material_type,
            raw_status=occurrence.raw_status,
        )
        existing = grouped.get(key)
        if existing is None:
            grouped[key] = SourceCandidate(
                key=key,
                candidate=candidate,
                provenance=(provenance,),
                identity_variants=(identity_variant,),
            )
            continue
        grouped[key] = SourceCandidate(
            key=key,
            candidate=existing.candidate,
            provenance=existing.provenance + (provenance,),
            identity_variants=existing.identity_variants + (identity_variant,),
        )
        duplicate_collapsed_count += 1

    ordered_candidates = tuple(grouped.values())
    ordered_rejections = tuple(rejected_rows)
    source_hash = sha256_text(
        _canonical_json(
            {
                "source_name": source_name,
                "source_row_count": raw_row_count,
                "candidates": [
                    {
                        "key": item.key,
                        "source_fingerprints": [entry.source_fingerprint for entry in item.provenance],
                    }
                    for item in ordered_candidates
                ],
                "rejected_rows": [
                    {
                        "source_row_id": item.source_row_id,
                        "reason_code": item.reason_code,
                        "source_fingerprint": item.source_fingerprint,
                    }
                    for item in ordered_rejections
                ],
            }
        )
    )
    return SourceSnapshot(
        source_name=source_name,
        source_hash=source_hash,
        source_row_count=raw_row_count,
        candidates=ordered_candidates,
        rejected_rows=ordered_rejections,
        duplicate_collapsed_count=duplicate_collapsed_count,
    )


def read_registry1(path: Path) -> SourceSnapshot:
    """Read the supplied Registry 1 workbook without contacting any source API."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    candidates: list[_CandidateOccurrence] = []
    rejected: list[RejectedSourceRow] = []
    raw_row_count = 0
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers or not any(_text(value) for value in headers):
                continue
            fallback_direction = _SHEET_DIRECTIONS.get(_header_key(sheet.title))
            for ordinal, row in enumerate(rows, start=2):
                if not any(_text(value) for value in row):
                    continue
                raw_row_count += 1
                source_row_id = f"registry1:{sheet.title}:{ordinal}"
                values = _field_values(headers, row)
                occurrence = _build_candidate(
                    source_name="registry1",
                    source_row_id=source_row_id,
                    values=values,
                    fallback_direction=fallback_direction,
                )
                candidate = occurrence.candidate
                if _source_key(candidate) is None:
                    rejected.append(
                        RejectedSourceRow(
                            source_name="registry1",
                            source_row_id=source_row_id,
                            reason_code="MISSING_IDENTITY",
                            source_fingerprint=candidate.source_fingerprint,
                        )
                    )
                else:
                    candidates.append(occurrence)
    finally:
        workbook.close()
    return _snapshot_from_candidates("registry1", raw_row_count, candidates, rejected)


def _read_registry2_rows(path: Path) -> list[tuple[str, dict[str, str]]]:
    if path.suffix.casefold() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload.get("rows", payload) if isinstance(payload, dict) else payload
        if not isinstance(rows, list) or not all(isinstance(row, Mapping) for row in rows):
            raise ValueError("REGISTRY2_SNAPSHOT_INVALID")
        return [
            (f"registry2:{ordinal}", {str(key): _text(value) for key, value in row.items()})
            for ordinal, row in enumerate(rows, start=1)
        ]
    with path.open(encoding="utf-8-sig", newline="") as source:
        return [
            (f"registry2:{ordinal}", {str(key): _text(value) for key, value in row.items()})
            for ordinal, row in enumerate(csv.DictReader(source), start=2)
        ]


def read_registry2_csv(path: Path) -> SourceSnapshot:
    """Read a captured Registry 2 CSV or JSON snapshot, never a live Sheet."""

    candidates: list[_CandidateOccurrence] = []
    rejected: list[RejectedSourceRow] = []
    rows = _read_registry2_rows(path)
    for source_row_id, raw_row in rows:
        values = _field_values(tuple(raw_row), tuple(raw_row.values()))
        occurrence = _build_candidate(
            source_name="registry2", source_row_id=source_row_id, values=values
        )
        candidate = occurrence.candidate
        if _source_key(candidate) is None:
            rejected.append(
                RejectedSourceRow(
                    source_name="registry2",
                    source_row_id=source_row_id,
                    reason_code="MISSING_IDENTITY",
                    source_fingerprint=candidate.source_fingerprint,
                )
            )
        else:
            candidates.append(occurrence)
    return _snapshot_from_candidates("registry2", len(rows), candidates, rejected)


def read_canonical_catalog(rows: Iterable[CanonicalClassification]) -> SourceSnapshot:
    """Turn the canonical MySQL catalog returned by the repository into a snapshot."""

    candidates: list[SourceCandidate] = []
    for row in rows:
        source_row_id = f"canonical_catalog:{row.content_entity_id}"
        payload = {
            "source_name": "canonical_catalog",
            "source_row_id": source_row_id,
            "title": normalize_title(row.title),
            "url": normalize_url(row.url).value,
            "material_id": None,
            "direction_code": row.direction_code,
            "material_type_code": row.material_type_code,
            "access_code": row.access_code,
            "lifecycle_code": row.lifecycle_code,
        }
        candidate = MaterialCandidate(**payload, source_fingerprint=sha256_text(_canonical_json(payload)))
        candidates.append(
            SourceCandidate(
                key=f"canonical:{row.content_entity_id}",
                candidate=candidate,
                provenance=(
                    SourceProvenance(
                        source_name="canonical_catalog",
                        source_row_id=source_row_id,
                        source_fingerprint=candidate.source_fingerprint,
                    ),
                ),
                identity_variants=(
                    SourceIdentityVariant(
                        source_row_id=source_row_id,
                        material_id=None,
                        normalized_url=candidate.url,
                        normalized_title=candidate.title,
                        material_type_code=candidate.material_type_code,
                    ),
                ),
            )
        )
    source_hash = sha256_text(
        _canonical_json(
            {
                "source_name": "canonical_catalog",
                "candidates": [
                    {"key": item.key, "source_fingerprint": item.source_fingerprint}
                    for item in candidates
                ],
            }
        )
    )
    return SourceSnapshot(
        source_name="canonical_catalog",
        source_hash=source_hash,
        source_row_count=len(candidates),
        candidates=tuple(candidates),
        rejected_rows=(),
    )
