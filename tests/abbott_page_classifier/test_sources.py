"""Captured Abbott registry source-reader contracts."""

from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from agents.abbott_page_classifier.domain import CanonicalClassification
from agents.abbott_page_classifier.sources import (
    read_canonical_catalog,
    read_registry1,
    read_registry2_csv,
)


FIXTURE_CSV = (
    Path(__file__).resolve().parents[1]
    / "fixtures"
    / "abbott_registry2_accepted_minimal.csv"
)


class RegistrySourceReaderTests(unittest.TestCase):
    def test_registry1_collapses_duplicate_material_and_preserves_provenance(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            fixture = Path(temporary_directory) / "registry1.xlsx"
            workbook = Workbook()
            cardiology = workbook.active
            cardiology.title = "кардио"
            dermatology = workbook.create_sheet("дерматология")
            headers = (
                "ID",
                "Название",
                "ссылка",
                "Направление",
                "доступ",
                "Тип контента",
            )
            for sheet in (cardiology, dermatology):
                sheet.append(headers)
            cardiology.append(
                (100, "Повторяемая КР", "https://abbottpro.ru/cardio/kr-100/", "Кардиология [262338]", "Врачи", "КР")
            )
            cardiology.append(
                (101, "Научная брошюра", "https://abbottpro.ru/cardio/brochure-101/", "Кардиология [262338]", "фарм", "Научно-брошюры")
            )
            cardiology.append(
                (102, "Материал без URL", "", "Кардиология [262338]", "Врачи", "Статьи")
            )
            dermatology.append(
                (100, "Повторяемая КР", "https://abbottpro.ru/cardio/kr-100/?utm_source=sheet", "Кардиология [262338]", "Врачи", "КР")
            )
            dermatology.append(
                (103, "Новая дерматология", "https://abbottpro.ru/dermatology/new-103/", "Дерматология [624635]", "Врачи", "Статьи")
            )
            dermatology.append(
                (104, "Новая таблица", "https://abbottpro.ru/dermatology/table-104/", "Дерматология [624635]", "Врачи", "Таблицы")
            )
            workbook.save(fixture)

            snapshot = read_registry1(fixture)
            rerun = read_registry1(fixture)

        self.assertEqual(snapshot.source_name, "registry1")
        self.assertEqual(snapshot.source_row_count, 6)
        self.assertEqual(len(snapshot.candidates), 5)
        self.assertEqual(snapshot.candidates_by_key["material:100"].material_type_code, "clinical_guidelines")
        self.assertEqual(len(snapshot.candidates_by_key["material:100"].provenance), 2)
        self.assertEqual(
            tuple(item.source_row_id for item in snapshot.candidates_by_key["material:100"].provenance),
            ("registry1:кардио:2", "registry1:дерматология:2"),
        )
        self.assertEqual(snapshot.candidates_by_key["material:101"].access_code, "pharmacists")
        self.assertEqual(snapshot.candidates_by_key["material:102"].url, "")
        self.assertEqual(snapshot.source_hash, rerun.source_hash)
        self.assertEqual(snapshot.rejected_rows, ())
        self.assertEqual(snapshot.outcome_count, snapshot.source_row_count)

    def test_duplicate_material_retains_every_normalized_identity_variant(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            fixture = Path(temporary_directory) / "registry1.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "кардио"
            sheet.append(("ID", "Название", "ссылка", "Тип контента"))
            sheet.append((100, "Первый", "HTTPS://ABBOTTPRO.RU/cardio/a/?utm_source=first", "Статьи"))
            sheet.append((100, "Второй", "https://abbottpro.ru/cardio/b/", "Статьи"))
            workbook.save(fixture)

            collapsed = read_registry1(fixture).candidates_by_key["material:100"]

        self.assertEqual(collapsed.material_id, "100")
        self.assertEqual(collapsed.url, "https://abbottpro.ru/cardio/a")
        self.assertEqual(len(collapsed.provenance), 2)
        self.assertEqual(
            tuple(
                (
                    variant.material_id,
                    variant.normalized_url,
                    variant.normalized_title,
                    variant.material_type_code,
                )
                for variant in collapsed.identity_variants
            ),
            (
                ("100", "https://abbottpro.ru/cardio/a", "Первый", "articles"),
                ("100", "https://abbottpro.ru/cardio/b", "Второй", "articles"),
            ),
        )

    def test_title_type_source_keys_keep_cyrillic_yo_distinct(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            fixture = Path(temporary_directory) / "registry2.csv"
            fixture.write_text(
                "Название,Тип контента\nВсе,Статьи\nВсё,Статьи\n",
                encoding="utf-8",
            )

            snapshot = read_registry2_csv(fixture)

        self.assertEqual(snapshot.source_row_count, 2)
        self.assertEqual(len(snapshot.candidates), 2)
        self.assertEqual(snapshot.duplicate_collapsed_count, 0)
        self.assertEqual({item.title for item in snapshot.candidates}, {"Все", "Всё"})

    def test_registry2_reads_only_captured_csv_and_rejects_rows_without_identity(self):
        snapshot = read_registry2_csv(FIXTURE_CSV)

        self.assertEqual(snapshot.source_name, "registry2")
        self.assertEqual(snapshot.source_row_count, 3)
        self.assertEqual(len(snapshot.candidates), 2)
        self.assertEqual(snapshot.candidates_by_key["material:200"].material_type_code, "educational_brochures")
        self.assertEqual(snapshot.rejected_rows[0].reason_code, "MISSING_IDENTITY")
        self.assertEqual(snapshot.outcome_count, snapshot.source_row_count)
        self.assertEqual(len(snapshot.source_hash), 64)

    def test_registry2_retains_raw_material_type_and_page_status_per_occurrence(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            fixture = Path(temporary_directory) / "registry2.csv"
            fixture.write_text(
                "ID,Название,URL,Направление,Тип материала,page_status\n"
                "200,Первый,https://abbottpro.ru/a,Кардиология,Архив,active\n"
                "200,Второй,https://abbottpro.ru/b,Кардиология,Статьи,Архив\n",
                encoding="utf-8",
            )

            collapsed = read_registry2_csv(fixture).candidates_by_key["material:200"]

        self.assertEqual(
            tuple(
                (variant.raw_material_type, variant.raw_status)
                for variant in collapsed.identity_variants
            ),
            (("Архив", "active"), ("Статьи", "Архив")),
        )

    def test_registry2_page_status_alias_maps_archive_lifecycle(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            fixture = Path(temporary_directory) / "registry2.csv"
            fixture.write_text(
                "ID,Название,URL,Направление,Тип контента,page_status\n"
                "201,Архивная статья,https://abbottpro.ru/archive,"
                "Кардиология,Статьи,Архив\n",
                encoding="utf-8",
            )

            item = read_registry2_csv(fixture).candidates_by_key["material:201"]

        self.assertEqual(item.lifecycle_code, "archive_candidate")
        self.assertEqual(item.identity_variants[0].raw_status, "Архив")

    def test_registry2_source_hash_binds_raw_archive_evidence(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            blank_fixture = Path(temporary_directory) / "blank.csv"
            archive_fixture = Path(temporary_directory) / "archive.csv"
            header = "ID,Название,URL,Направление,Тип материала,page_status\n"
            blank_fixture.write_text(
                header
                + "202,Материал,https://abbottpro.ru/material,Кардиология,,active\n",
                encoding="utf-8",
            )
            archive_fixture.write_text(
                header
                + "202,Материал,https://abbottpro.ru/material,"
                "Кардиология,Архив,active\n",
                encoding="utf-8",
            )

            blank = read_registry2_csv(blank_fixture)
            archive = read_registry2_csv(archive_fixture)

        self.assertNotEqual(blank.source_hash, archive.source_hash)

    def test_canonical_catalog_rows_have_deterministic_source_identity(self):
        rows = (
            CanonicalClassification(
                content_entity_id=41,
                title="Canonical material",
                url="https://abbottpro.ru/cardio/canonical/",
                direction_code="cardiology",
                material_type_code="articles",
                access_code="doctors",
                lifecycle_code="active",
            ),
        )

        snapshot = read_canonical_catalog(rows)

        self.assertEqual(snapshot.source_name, "canonical_catalog")
        self.assertEqual(snapshot.source_row_count, 1)
        self.assertEqual(snapshot.candidates_by_key["canonical:41"].source_row_id, "canonical_catalog:41")
        self.assertEqual(snapshot.outcome_count, 1)


if __name__ == "__main__":
    unittest.main()
