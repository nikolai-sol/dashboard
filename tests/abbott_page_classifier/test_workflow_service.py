"""Executable fake integration for the canonical weekly proposal service."""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path
import json
import tempfile
import unittest
from datetime import date
from unittest.mock import patch

from openpyxl import Workbook

from agents.abbott_page_classifier.approval_hashes import compute_taxonomy_digest
from agents.abbott_page_classifier.batch_service import PersistedApprovalBatch
from agents.abbott_page_classifier.domain import (
    CanonicalClassification,
    Proposal,
    TAXONOMY_LABELS,
    TaxonomyVersion,
)
from agents.abbott_page_classifier.identity import IdentityAlias, IdentityResolver
from agents.abbott_page_classifier.llm_classifier import (
    LLM_PRIMARY_MODEL,
    LlmAttempt,
    LlmClassification,
)
from agents.abbott_page_classifier.normalization import sha256_text
from agents.abbott_page_classifier.reconcile import reconcile_entity
from agents.abbott_page_classifier.workflow_service import (
    CanonicalWeeklyProposalService,
    ObservedPage,
    ReconciliationContext,
    WorkflowConfiguration,
)
from agents.abbott_page_classifier.workflow_repository import _input_payload
from agents.abbott_page_classifier.sources import collapse_observed_pages


TERMS = {kind: tuple(sorted(labels)) for kind, labels in TAXONOMY_LABELS.items()}
TAXONOMY = TaxonomyVersion(
    version="abbott.v1",
    terms=TERMS,
    digest=compute_taxonomy_digest("abbott.v1", TERMS),
)
CONFIG = WorkflowConfiguration(
    taxonomy_version="abbott.v1",
    prompt_version="prompt.v1",
    model_routing_version="routing.v1",
    code_revision="a" * 40,
)


class StatefulWorkflowStore:
    """Stateful fake exercising the same cross-process service protocol."""

    def __init__(self, context):
        self.context = context
        self.runs_by_id = {}
        self.run_id_by_key = {}
        self.batches_by_run = {}
        self.attempts = []
        self.entity_creations = []
        self.entity_by_run_item = {}
        self.persisted_attempts = {}
        self.next_entity_id = 100

    def load_reconciliation_context(self, configuration):
        self.asserted_configuration = configuration
        return self.context

    def persist_reconciliation_run(self, draft):
        existing = self.run_id_by_key.get(draft.run_key)
        if existing is not None:
            return self.runs_by_id[existing]
        run = replace(draft, run_id=len(self.runs_by_id) + 1)
        self.runs_by_id[run.run_id] = run
        self.run_id_by_key[run.run_key] = run.run_id
        return run

    def load_reconciliation_run(self, run_id):
        return self.runs_by_id[int(run_id)]

    def load_finalized_batch_for_run(self, run_id):
        return self.batches_by_run.get(int(run_id))

    def load_llm_attempts(self, run_id, item_key):
        return dict(self.persisted_attempts.get((int(run_id), item_key), {}))

    def resolve_or_create_registry1_entities(self, run_id, item_keys):
        resolved = {}
        run = self.load_reconciliation_run(run_id)
        by_key = {item.item_key: item for item in run.items}
        for item_key in item_keys:
            item = by_key[item_key]
            self.assertEqual(item.identity_status, "new")
            self.assertIsNotNone(item.reconciliation_input.registry1)
            existing = self.entity_by_run_item.get((int(run_id), item_key))
            if existing is not None:
                resolved[item_key] = existing
                continue
            self.next_entity_id += 1
            resolved[item_key] = self.next_entity_id
            self.entity_by_run_item[(int(run_id), item_key)] = self.next_entity_id
            self.entity_creations.append((run_id, item_key, self.next_entity_id))
        return resolved

    def append_llm_attempt(self, run_id, item_key, route_kind, attempt):
        self.attempts.append((int(run_id), item_key, route_kind, attempt))
        self.persisted_attempts.setdefault((int(run_id), item_key), {})[route_kind] = attempt

    def finalize_reconciliation_run(self, run_id, batch):
        existing = self.batches_by_run.get(int(run_id))
        if existing is not None:
            self.assertEqual(existing.batch, batch)
            return existing
        persisted = PersistedApprovalBatch(batch=batch, database_batch_id=70 + int(run_id))
        self.batches_by_run[int(run_id)] = persisted
        self.runs_by_id[int(run_id)] = replace(
            self.runs_by_id[int(run_id)], status="finalized"
        )
        return persisted

    # unittest-style assertions keep fake violations executable, not flags.
    def assertEqual(self, left, right):
        if left != right:
            raise AssertionError((left, right))

    def assertIsNotNone(self, value):
        if value is None:
            raise AssertionError("expected value")


class RecordingClassifier:
    def __init__(self):
        self.requests = []

    def classify(self, request, model):
        self.requests.append((request, model))
        return LlmAttempt(
            status="success",
            model=model,
            classification=LlmClassification(
                direction_code="cardiology" if "direction_code" in request.requested_fields else None,
                material_type_code="articles" if "material_type_code" in request.requested_fields else None,
                direction_confidence=0.99 if "direction_code" in request.requested_fields else None,
                material_type_confidence=0.99 if "material_type_code" in request.requested_fields else None,
                alternative_direction_codes=(),
                evidence=("bounded source metadata",),
                requires_medical_review=False,
                insufficient_evidence=False,
            ),
            attempt_count=1,
            input_hash="f" * 64,
            requested_fields=request.requested_fields,
            elapsed_ms=3,
        )


def context(
    *,
    entities=(),
    aliases=(),
    taxonomy=TAXONOMY,
    predecessor_content_entity_ids=(),
    predecessor_catalog_entities=None,
    observed_pages=(),
):
    if predecessor_catalog_entities is None:
        predecessor_ids = set(predecessor_content_entity_ids)
        predecessor_catalog_entities = tuple(
            entity
            for entity in entities
            if entity.content_entity_id in predecessor_ids
        )
    return ReconciliationContext(
        predecessor_release_id=8,
        predecessor_snapshot_ids=(11, 12),
        predecessor_snapshot_digests=("1" * 64, "2" * 64),
        taxonomy=taxonomy,
        entities=tuple(entities),
        aliases=tuple(aliases),
        predecessor_content_entity_ids=tuple(
            predecessor_content_entity_ids
        ),
        predecessor_catalog_entities=tuple(predecessor_catalog_entities),
        observed_pages=tuple(observed_pages),
    )


def write_sources(
    directory: Path,
    *,
    material_id="900",
    direction="Кардиология",
    material_type="Статьи",
    registry2_material_id="",
    registry2_url="",
    registry2_direction="",
    registry2_material_type="",
):
    registry1 = directory / "registry1.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "кардио"
    sheet.append(("ID", "Название", "ссылка", "Направление", "Тип контента"))
    sheet.append((material_id, "Новый материал", "https://abbottpro.ru/cardio/new", direction, material_type))
    workbook.save(registry1)
    registry2 = directory / "registry2.csv"
    rows = "ID,Название,URL,Направление,Тип материала\n"
    if registry2_material_id or registry2_url:
        rows += (
            f"{registry2_material_id},Новый материал,{registry2_url},"
            f"{registry2_direction},{registry2_material_type}\n"
        )
    registry2.write_text(rows, encoding="utf-8")
    return registry1, registry2


def write_empty_sources(directory: Path):
    registry1 = directory / "registry1.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "кардио"
    sheet.append(("ID", "Название", "ссылка", "Направление", "Тип контента"))
    workbook.save(registry1)
    registry2 = directory / "registry2.csv"
    registry2.write_text(
        "ID,Название,URL,Направление,Тип материала\n", encoding="utf-8"
    )
    return registry1, registry2


class WeeklyProposalServiceTests(unittest.TestCase):
    def test_unaliased_observed_page_creates_one_immutable_review_item(self):
        observed = ObservedPage(
            "https://abbottpro.ru/auth", "Вход", 7, date(2026, 8, 1), date(2026, 8, 2)
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(context(observed_pages=(observed,)))
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(registry1, registry2)

        item = store.runs_by_id[receipt.run_id].items[0]
        self.assertEqual(receipt.catalog_gap_count, 1)
        self.assertEqual(item.identity_status, "new")
        self.assertEqual(item.reconciliation_input.registry1.url, observed.normalized_url)
        self.assertEqual(item.reconciliation_input.deterministic_proposal.rule_code, "SERVICE_ROUTE")
        self.assertIsNone(
            item.reconciliation_input.deterministic_proposal.material_type_code
        )
        self.assertIsInstance(
            item.reconciliation_input.deterministic_proposal.confidence, float
        )
        self.assertEqual(
            reconcile_entity(item.reconciliation_input).readiness_state,
            "unresolved",
        )

    def test_repeated_observed_page_rows_collapse_before_item_creation(self):
        first = ObservedPage("https://abbottpro.ru/unknown", "One", 2, date(2026, 8, 1), date(2026, 8, 1))
        second = ObservedPage("https://abbottpro.ru/unknown", "Two", 3, date(2026, 8, 2), date(2026, 8, 2))
        collapsed = collapse_observed_pages((first, second))
        self.assertEqual(
            collapsed,
            (
                ObservedPage(
                    "https://abbottpro.ru/unknown",
                    "Two",
                    5,
                    date(2026, 8, 1),
                    date(2026, 8, 2),
                    (sha256_text("https://abbottpro.ru/unknown"),),
                ),
            ),
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(context(observed_pages=(first, second)))
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(registry1, registry2)

        self.assertEqual(receipt.catalog_gap_count, 1)
        self.assertEqual(len(store.runs_by_id[receipt.run_id].items), 1)

    def test_observed_query_variants_collapse_to_one_page_path_identity(self):
        first = ObservedPage(
            "https://abbottpro.ru/academy/video?session=one",
            "Видео",
            2,
            date(2026, 8, 1),
            date(2026, 8, 1),
        )
        second = ObservedPage(
            "https://abbottpro.ru/academy/video?session=two&utm_source=email",
            "Видео",
            3,
            date(2026, 8, 2),
            date(2026, 8, 2),
        )

        collapsed = collapse_observed_pages((first, second))

        self.assertEqual(
            collapsed,
            (
                ObservedPage(
                    "https://abbottpro.ru/academy/video",
                    "Видео",
                    5,
                    date(2026, 8, 1),
                    date(2026, 8, 2),
                    tuple(
                        sorted(
                            (
                                sha256_text(
                                    "https://abbottpro.ru/academy/video?session=one"
                                ),
                                sha256_text(
                                    "https://abbottpro.ru/academy/video?session=two"
                                ),
                            )
                        )
                    ),
                ),
            ),
        )

    def test_observed_unknown_query_variant_stays_query_free_and_unresolved(self):
        entity = CanonicalClassification(
            7,
            "Known",
            "https://abbottpro.ru/academy/known",
            "cardiology",
            "video",
            "all",
            "active",
            1,
        )
        observed = ObservedPage(
            "https://abbottpro.ru/academy/known?session=one",
            "Known",
            3,
            date(2026, 8, 1),
            date(2026, 8, 1),
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(
                context(entities=(entity,), observed_pages=(observed,))
            )
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )

        self.assertEqual(receipt.catalog_gap_count, 1)
        item = store.runs_by_id[receipt.run_id].items[0]
        self.assertEqual(
            item.reconciliation_input.registry1.url,
            "https://abbottpro.ru/academy/known",
        )
        self.assertNotIn("?", item.grouping_key)
        self.assertEqual(item.identity_status, "new")
        durable_payload = json.dumps(
            _input_payload(item.reconciliation_input),
            ensure_ascii=False,
            sort_keys=True,
        )
        self.assertNotIn("?", durable_payload)
        self.assertNotIn("session=one", durable_payload)

    def test_observed_variants_all_bound_to_one_exact_target_are_skipped(self):
        entity = CanonicalClassification(
            7,
            "Known",
            "https://abbottpro.ru/academy/known?session=one",
            "cardiology",
            "video",
            "all",
            "active",
            1,
        )
        observed = (
            ObservedPage(
                "https://abbottpro.ru/academy/known?session=one",
                "Known",
                2,
                date(2026, 8, 1),
                date(2026, 8, 1),
            ),
            ObservedPage(
                "https://abbottpro.ru/academy/known?session=two",
                "Known",
                3,
                date(2026, 8, 2),
                date(2026, 8, 2),
            ),
        )
        aliases = (
            IdentityAlias(
                7,
                "url",
                "https://abbottpro.ru/academy/known?session=two",
                "strong",
            ),
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(
                context(entities=(entity,), aliases=aliases, observed_pages=observed)
            )
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )

        self.assertEqual(receipt.catalog_gap_count, 0)
        self.assertEqual(store.runs_by_id[receipt.run_id].items, ())

    def test_observed_variants_bound_to_distinct_targets_create_one_collision(self):
        first_entity = CanonicalClassification(
            7, "First", "https://abbottpro.ru/shared?version=one",
            "cardiology", "articles", "all", "active", 1,
        )
        second_entity = CanonicalClassification(
            8, "Second", "https://abbottpro.ru/shared?version=two",
            "gastroenterology", "articles", "all", "active", 2,
        )
        observed = (
            ObservedPage(
                first_entity.url, "First", 2, date(2026, 8, 1), date(2026, 8, 1)
            ),
            ObservedPage(
                second_entity.url, "Second", 3, date(2026, 8, 2), date(2026, 8, 2)
            ),
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(
                context(entities=(first_entity, second_entity), observed_pages=observed)
            )
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )

        self.assertEqual(receipt.catalog_gap_count, 1)
        item = store.runs_by_id[receipt.run_id].items[0]
        self.assertEqual(item.identity_status, "collision")
        self.assertNotIn("?", item.grouping_key)

    def test_known_observed_alias_does_not_create_duplicate_item(self):
        entity = CanonicalClassification(7, "Known", "https://abbottpro.ru/known", "cardiology", "articles", "all", "active", 1)
        observed = ObservedPage("https://abbottpro.ru/known", "Known", 3, date(2026, 8, 1), date(2026, 8, 1))
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(context(entities=(entity,), aliases=(IdentityAlias(7, "url", entity.url, "strong"),), observed_pages=(observed,)))
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(registry1, registry2)

        self.assertEqual(receipt.catalog_gap_count, 0)
        self.assertEqual(store.runs_by_id[receipt.run_id].items, ())

    def test_known_observed_alias_without_metadata_creates_matched_review_item(self):
        entity = CanonicalClassification(
            7,
            "Known without metadata",
            "https://abbottpro.ru/known",
            None,
            None,
            None,
            "unknown",
            None,
        )
        observed = ObservedPage(
            entity.url,
            entity.title,
            3,
            date(2026, 8, 1),
            date(2026, 8, 1),
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(context(
                entities=(entity,),
                aliases=(IdentityAlias(7, "url", entity.url, "strong"),),
                observed_pages=(observed,),
            ))
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )

        self.assertEqual(receipt.catalog_gap_count, 1)
        item = store.runs_by_id[receipt.run_id].items[0]
        self.assertEqual(item.identity_status, "matched")
        self.assertEqual(item.content_entity_id, 7)
        self.assertEqual(item.reconciliation_input.active_canonical, entity)
        reconciled = reconcile_entity(item.reconciliation_input)
        self.assertEqual(reconciled.final_access_code, "unspecified")
        self.assertEqual(reconciled.final_lifecycle_code, "active")

    def test_unique_weak_observed_slug_match_still_creates_review_item(self):
        entity = CanonicalClassification(7, "Known", "https://abbottpro.ru/cardio/known", "cardiology", "articles", "all", "active", 1)
        observed = ObservedPage("https://abbottpro.ru/cardio/known?version=2", "Other title", 3, date(2026, 8, 1), date(2026, 8, 1))
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(context(entities=(entity,), aliases=(IdentityAlias(7, "slug", "known", "weak"),), observed_pages=(observed,)))
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(registry1, registry2)

        self.assertEqual(receipt.catalog_gap_count, 1)
        self.assertEqual(store.runs_by_id[receipt.run_id].items[0].identity_status, "new")

    def test_observed_strong_url_collision_is_a_collision_item(self):
        left = CanonicalClassification(7, "Left", "https://abbottpro.ru/left", "cardiology", "articles", "all", "active", 1)
        right = CanonicalClassification(8, "Right", "https://abbottpro.ru/right", "cardiology", "articles", "all", "active", 2)
        observed = ObservedPage("https://abbottpro.ru/collision", "Collision", 3, date(2026, 8, 1), date(2026, 8, 1))
        aliases = (IdentityAlias(7, "url", observed.normalized_url, "strong"), IdentityAlias(8, "url", observed.normalized_url, "strong"))
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(context(entities=(left, right), aliases=aliases, observed_pages=(observed,)))
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(registry1, registry2)

        item = store.runs_by_id[receipt.run_id].items[0]
        self.assertEqual(item.identity_status, "collision")
        self.assertTrue(item.reconciliation_input.identity_conflict)

    def test_observed_pages_use_preindexed_strong_urls(self):
        observed = (
            ObservedPage("https://abbottpro.ru/one", "One", 1, date(2026, 8, 1), date(2026, 8, 1)),
            ObservedPage("https://abbottpro.ru/two", "Two", 1, date(2026, 8, 1), date(2026, 8, 1)),
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(context(observed_pages=observed))
            with patch.object(
                IdentityResolver,
                "resolve",
                side_effect=AssertionError("observed URL resolution must use the prebuilt index"),
            ):
                receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                    registry1, registry2
                )

        self.assertEqual(receipt.catalog_gap_count, 2)

    def test_registry_rows_use_prepared_identity_index(self):
        active = CanonicalClassification(
            7,
            "Known",
            "https://abbottpro.ru/cardio/new",
            "cardiology",
            "articles",
            "all",
            "active",
            1,
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(Path(temporary))
            store = StatefulWorkflowStore(context(entities=(active,)))
            with patch.object(
                IdentityResolver,
                "resolve",
                side_effect=AssertionError(
                    "registry identity resolution must use the prepared index"
                ),
            ):
                receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                    registry1, registry2
                )

        self.assertEqual(receipt.source_count, 1)
        self.assertEqual(store.runs_by_id[receipt.run_id].items[0].identity_status, "matched")

    def test_active_catalog_gap_is_included_and_classified_without_entity_creation(self):
        entity = CanonicalClassification(
            7,
            "Материал без направления",
            "",
            None,
            "articles",
            "all",
            "active",
            17,
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            classifier = RecordingClassifier()
            store = StatefulWorkflowStore(
                context(
                    entities=(entity,),
                    predecessor_content_entity_ids=(entity.content_entity_id,),
                    predecessor_catalog_entities=(
                        replace(
                            entity,
                            url="https://abbottpro.ru/articles/gap",
                        ),
                    ),
                )
            )
            service = CanonicalWeeklyProposalService(
                store, CONFIG, classifier_factory=lambda: classifier
            )
            first = service.reconcile(registry1, registry2)
            replay = service.reconcile(registry1, registry2)
            receipt = service.classify(first.run_id, execute_llm=True)

        self.assertEqual(first.run_id, replay.run_id)
        self.assertEqual(first.source_count, 0)
        self.assertEqual(first.catalog_gap_count, 1)
        persisted = store.load_reconciliation_run(first.run_id)
        self.assertEqual(len(persisted.items), 1)
        gap = persisted.items[0]
        self.assertEqual(gap.identity_status, "matched")
        self.assertEqual(gap.content_entity_id, entity.content_entity_id)
        self.assertEqual(gap.reconciliation_input.registry1.source_name, "canonical_catalog")
        self.assertEqual(
            gap.reconciliation_input.active_canonical.url,
            "https://abbottpro.ru/articles/gap",
        )
        self.assertEqual(store.entity_creations, [])
        self.assertEqual(receipt.eligible_count, 1)
        self.assertEqual(receipt.ready_count, 1)
        self.assertEqual(len(classifier.requests), 1)
        self.assertEqual(classifier.requests[0][0].requested_fields, ("direction_code",))

    def test_catalog_gap_includes_archive_candidate_but_not_archived_entity(self):
        entities = (
            CanonicalClassification(
                7, "Complete", "https://abbottpro.ru/complete", "cardiology",
                "articles", "all", "active", 17,
            ),
            CanonicalClassification(
                8, "Archived", "https://abbottpro.ru/archived", None,
                "articles", "all", "archived", 18,
            ),
            CanonicalClassification(
                9, "Archive candidate", "https://abbottpro.ru/archive-candidate", None,
                "articles", "all", "archive_candidate", 19,
            ),
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(
                context(
                    entities=entities,
                    predecessor_content_entity_ids=(7, 8, 9),
                )
            )
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )

        self.assertEqual(receipt.catalog_gap_count, 1)
        items = store.load_reconciliation_run(receipt.run_id).items
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0].content_entity_id, 9)

    def test_catalog_gap_does_not_duplicate_entity_already_in_registry_source(self):
        entity = CanonicalClassification(
            7,
            "Новый материал",
            "https://abbottpro.ru/cardio/new",
            None,
            "articles",
            "all",
            "active",
            17,
        )
        aliases = (IdentityAlias(7, "material_id", "900", "strong"),)
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(
                Path(temporary), direction="", material_type="Статьи"
            )
            store = StatefulWorkflowStore(
                context(
                    entities=(entity,),
                    aliases=aliases,
                    predecessor_content_entity_ids=(7,),
                )
            )
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )

        persisted = store.load_reconciliation_run(receipt.run_id)
        self.assertEqual(receipt.catalog_gap_count, 0)
        self.assertEqual(len(persisted.items), 1)
        self.assertEqual(persisted.items[0].content_entity_id, 7)
        self.assertEqual(persisted.items[0].reconciliation_input.registry1.source_name, "registry1")

    def test_incomplete_registry_entity_outside_active_predecessor_is_not_a_gap(self):
        entity = CanonicalClassification(
            7, "Old", "https://abbottpro.ru/old", None,
            "articles", "all", "active", 17,
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            store = StatefulWorkflowStore(context(entities=(entity,)))
            receipt = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )

        self.assertEqual(receipt.catalog_gap_count, 0)
        self.assertEqual(store.load_reconciliation_run(receipt.run_id).items, ())

    def test_identical_new_strong_identity_merges_registry1_and_registry2(self):
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(
                Path(temporary),
                registry2_material_id="900",
                registry2_url="https://abbottpro.ru/cardio/new",
                registry2_direction="Кардиология",
                registry2_material_type="Статьи",
            )
            store = StatefulWorkflowStore(context())
            run = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )

        persisted = store.load_reconciliation_run(run.run_id)
        self.assertEqual(len(persisted.items), 1)
        item = persisted.items[0].reconciliation_input
        self.assertIsNotNone(item.registry1)
        self.assertIsNotNone(item.registry2)
        self.assertFalse(item.identity_conflict)

    def test_registry2_only_identity_is_rejected_without_entity_creation(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            registry1 = directory / "registry1.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "кардио"
            sheet.append(
                ("ID", "Название", "ссылка", "Направление", "Тип контента")
            )
            workbook.save(registry1)
            registry2 = directory / "registry2.csv"
            registry2.write_text(
                "ID,Название,URL,Направление,Тип материала\n"
                "900,Новый материал,https://abbottpro.ru/cardio/new,"
                "Кардиология,Статьи\n",
                encoding="utf-8",
            )
            store = StatefulWorkflowStore(context())
            service = CanonicalWeeklyProposalService(store, CONFIG)
            run = service.reconcile(registry1, registry2)
            replay = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )
            receipt = service.classify(run.run_id, execute_llm=True)

        persisted = store.load_reconciliation_run(run.run_id)
        self.assertEqual(run.run_id, replay.run_id)
        self.assertEqual(len(persisted.items), 1)
        source_item = persisted.items[0]
        self.assertEqual(source_item.identity_status, "rejected")
        self.assertIsNone(source_item.reconciliation_input.registry1)
        self.assertIsNotNone(source_item.reconciliation_input.registry2)
        self.assertEqual(
            source_item.reconciliation_input.rejection_code,
            "REGISTRY1_IDENTITY_REQUIRED",
        )
        self.assertEqual(receipt.ready_count, 0)
        self.assertEqual(receipt.rejected_count, 1)
        self.assertEqual(store.entity_creations, [])
        batch_item = store.batches_by_run[run.run_id].batch.items[0]
        self.assertEqual(batch_item.readiness_state, "rejected")
        self.assertEqual(
            batch_item.decision_reason, "REGISTRY1_IDENTITY_REQUIRED"
        )
        self.assertIn(
            "REGISTRY1_IDENTITY_REQUIRED",
            tuple(code.value for code in batch_item.conflict_codes),
        )
        self.assertIsNotNone(batch_item.registry2_values)

    def test_differing_new_strong_evidence_is_identity_collision(self):
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(
                Path(temporary),
                registry2_material_id="901",
                registry2_url="https://abbottpro.ru/cardio/new",
                registry2_direction="Кардиология",
                registry2_material_type="Статьи",
            )
            store = StatefulWorkflowStore(context())
            run = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )
            receipt = CanonicalWeeklyProposalService(store, CONFIG).classify(
                run.run_id, execute_llm=False
            )

        self.assertEqual(len(store.load_reconciliation_run(run.run_id).items), 1)
        self.assertEqual(receipt.conflict_count, 1)
        self.assertEqual(store.entity_creations, [])

    def test_reconciliation_is_content_addressed_and_repeat_safe(self):
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(Path(temporary))
            store = StatefulWorkflowStore(context())
            service = CanonicalWeeklyProposalService(store, CONFIG)
            first = service.reconcile(registry1, registry2)
            second = CanonicalWeeklyProposalService(store, CONFIG).reconcile(registry1, registry2)

        self.assertEqual(first.run_id, second.run_id)
        self.assertEqual(first.run_key, second.run_key)
        self.assertEqual(first.registry1_hash, second.registry1_hash)
        self.assertEqual(len(first.run_key), 64)
        run = store.load_reconciliation_run(first.run_id)
        self.assertEqual(run.context.predecessor_snapshot_ids, (11, 12))
        self.assertNotEqual(run.registry1.source_hash, run.context.predecessor_snapshot_digests[0])

    def test_reconciliation_key_changes_with_sanitized_observed_evidence(self):
        first_page = ObservedPage(
            "https://abbottpro.ru/unknown?session=one",
            "Unknown",
            2,
            date(2026, 8, 1),
            date(2026, 8, 1),
        )
        second_page = replace(first_page, pageviews=3)
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_empty_sources(Path(temporary))
            first = CanonicalWeeklyProposalService(
                StatefulWorkflowStore(context(observed_pages=(first_page,))), CONFIG
            ).reconcile(registry1, registry2)
            second = CanonicalWeeklyProposalService(
                StatefulWorkflowStore(context(observed_pages=(second_page,))), CONFIG
            ).reconcile(registry1, registry2)

        self.assertNotEqual(first.run_key, second.run_key)

    def test_two_missing_identity_rows_persist_as_distinct_deterministic_items(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            registry1 = directory / "registry1.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "кардио"
            sheet.append(("ID", "Название", "ссылка", "Направление", "Тип контента"))
            sheet.append(("", "", "", "Кардиология", "Статьи"))
            sheet.append(("", "", "", "Кардиология", "Статьи"))
            workbook.save(registry1)
            registry2 = directory / "registry2.csv"
            registry2.write_text(
                "ID,Название,URL,Направление,Тип материала\n", encoding="utf-8"
            )

            store = StatefulWorkflowStore(context())
            service = CanonicalWeeklyProposalService(store, CONFIG)
            first = service.reconcile(registry1, registry2)
            replay = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )
            persisted = store.load_reconciliation_run(first.run_id)

            independent_store = StatefulWorkflowStore(context())
            independent = CanonicalWeeklyProposalService(
                independent_store, CONFIG
            ).reconcile(registry1, registry2)
            rebuilt = independent_store.load_reconciliation_run(independent.run_id)

        self.assertEqual(first.run_id, replay.run_id)
        self.assertEqual(len(persisted.items), 2)
        self.assertEqual(
            {item.identity_status for item in persisted.items}, {"rejected"}
        )
        self.assertEqual(len({item.item_key for item in persisted.items}), 2)
        self.assertEqual(len({item.input_hash for item in persisted.items}), 2)
        self.assertEqual(
            tuple(
                item.reconciliation_input.rejected_source_row.source_row_id
                for item in persisted.items
            ),
            ("registry1:кардио:2", "registry1:кардио:3"),
        )
        self.assertEqual(
            tuple((item.item_key, item.input_hash) for item in persisted.items),
            tuple((item.item_key, item.input_hash) for item in rebuilt.items),
        )

    def test_classification_creates_registry1_entity_before_finalizing_ready_batch(self):
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(Path(temporary))
            store = StatefulWorkflowStore(context())
            first_process = CanonicalWeeklyProposalService(store, CONFIG)
            run = first_process.reconcile(registry1, registry2)
            second_process = CanonicalWeeklyProposalService(store, CONFIG)
            receipt = second_process.classify(run.run_id, execute_llm=False)

        self.assertEqual(receipt.batch_id, 71)
        self.assertEqual(len(store.entity_creations), 1)
        item = store.batches_by_run[run.run_id].batch.items[0]
        self.assertEqual(item.readiness_state, "ready")
        self.assertEqual(item.content_entity_id, 101)
        self.assertEqual(store.batches_by_run[run.run_id].batch.source_snapshot_ids, (11, 12))

    def test_strong_identity_collision_stays_non_ready_and_creates_no_entity(self):
        entities = (
            CanonicalClassification(1, "One", "https://abbottpro.ru/one", "cardiology", "articles", "all", "active", 10),
            CanonicalClassification(2, "Two", "https://abbottpro.ru/cardio/new", "cardiology", "articles", "all", "active", 20),
        )
        aliases = (IdentityAlias(1, "material_id", "900", "strong"),)
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(Path(temporary))
            store = StatefulWorkflowStore(context(entities=entities, aliases=aliases))
            service = CanonicalWeeklyProposalService(store, CONFIG)
            run = service.reconcile(registry1, registry2)
            receipt = service.classify(run.run_id, execute_llm=False)

        item = store.batches_by_run[run.run_id].batch.items[0]
        self.assertEqual(item.readiness_state, "conflict")
        self.assertIsNone(item.content_entity_id)
        self.assertEqual(store.entity_creations, [])
        self.assertEqual(receipt.conflict_count, 1)

    def test_eligible_fields_append_sanitized_primary_attempt_before_finalization(self):
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(Path(temporary), direction="", material_type="")
            store = StatefulWorkflowStore(context())
            classifier = RecordingClassifier()
            service = CanonicalWeeklyProposalService(store, CONFIG, classifier_factory=lambda: classifier)
            run = service.reconcile(registry1, registry2)
            receipt = service.classify(run.run_id, execute_llm=True)

        self.assertEqual(receipt.eligible_count, 1)
        self.assertEqual(len(classifier.requests), 1)
        self.assertEqual(classifier.requests[0][1], LLM_PRIMARY_MODEL)
        self.assertEqual(len(store.attempts), 1)
        self.assertEqual(store.attempts[0][2], "terra_primary")
        self.assertEqual(store.batches_by_run[run.run_id].batch.items[0].readiness_state, "ready")

    def test_repeat_classification_attests_the_single_batch_for_the_run(self):
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(Path(temporary))
            store = StatefulWorkflowStore(context())
            run = CanonicalWeeklyProposalService(store, CONFIG).reconcile(registry1, registry2)
            first = CanonicalWeeklyProposalService(store, CONFIG).classify(run.run_id, execute_llm=False)
            second = CanonicalWeeklyProposalService(store, CONFIG).classify(run.run_id, execute_llm=False)

        self.assertEqual(first.batch_id, second.batch_id)
        self.assertEqual(first.batch_key, second.batch_key)
        self.assertEqual(len(store.batches_by_run), 1)

    def test_finalized_replay_returns_before_constructing_or_calling_provider(self):
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(
                Path(temporary), direction="", material_type=""
            )
            store = StatefulWorkflowStore(context())
            run = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )
            first_classifier = RecordingClassifier()
            first = CanonicalWeeklyProposalService(
                store, CONFIG, classifier_factory=lambda: first_classifier
            ).classify(run.run_id, execute_llm=True)
            replay_classifier = RecordingClassifier()
            second = CanonicalWeeklyProposalService(
                store, CONFIG, classifier_factory=lambda: replay_classifier
            ).classify(run.run_id, execute_llm=True)

        self.assertEqual(first.batch_id, second.batch_id)
        self.assertEqual(replay_classifier.requests, [])

    def test_partial_llm_recovery_reuses_persisted_primary_attempt(self):
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(
                Path(temporary), direction="", material_type=""
            )
            store = StatefulWorkflowStore(context())
            run = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )
            item_key = store.load_reconciliation_run(run.run_id).items[0].item_key
            # This is the exact durable primary row left by an interrupted process.
            persisted = LlmAttempt(
                status="success",
                model=LLM_PRIMARY_MODEL,
                classification=LlmClassification(
                    direction_code=None,
                    material_type_code="articles",
                    direction_confidence=None,
                    material_type_confidence=0.99,
                    alternative_direction_codes=(),
                    evidence=("bounded source metadata",),
                    requires_medical_review=False,
                    insufficient_evidence=False,
                ),
                attempt_count=1,
                input_hash="f" * 64,
                requested_fields=("material_type_code",),
                elapsed_ms=3,
            )
            store.persisted_attempts[(run.run_id, item_key)] = {
                "terra_primary": persisted
            }
            classifier = RecordingClassifier()
            receipt = CanonicalWeeklyProposalService(
                store, CONFIG, classifier_factory=lambda: classifier
            ).classify(run.run_id, execute_llm=True)

        self.assertEqual(receipt.ready_count, 1)
        self.assertEqual(classifier.requests, [])

    def test_llm_disagreement_survives_routing_into_conflict_queue(self):
        class DisagreeingClassifier:
            def classify(self, request, model):
                material_type = "articles" if model == LLM_PRIMARY_MODEL else "video"
                return LlmAttempt(
                    status="success",
                    model=model,
                    classification=LlmClassification(
                        direction_code=None,
                        material_type_code=material_type,
                        direction_confidence=None,
                        material_type_confidence=(
                            0.80 if model == LLM_PRIMARY_MODEL else 0.99
                        ),
                        alternative_direction_codes=(),
                        evidence=("bounded source metadata",),
                        requires_medical_review=False,
                        insufficient_evidence=False,
                    ),
                    attempt_count=1,
                    input_hash="f" * 64,
                    requested_fields=request.requested_fields,
                    elapsed_ms=3,
                )

        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(
                Path(temporary), direction="", material_type=""
            )
            store = StatefulWorkflowStore(context())
            run = CanonicalWeeklyProposalService(store, CONFIG).reconcile(
                registry1, registry2
            )
            receipt = CanonicalWeeklyProposalService(
                store, CONFIG, classifier_factory=DisagreeingClassifier
            ).classify(run.run_id, execute_llm=True)

        item = store.batches_by_run[run.run_id].batch.items[0]
        self.assertEqual(receipt.conflict_count, 1)
        self.assertIn(
            "LLM_DISAGREEMENT",
            tuple(getattr(code, "value", str(code)) for code in item.conflict_codes),
        )

    def test_non_v1_reconciliation_classifies_in_a_fresh_service_process(self):
        taxonomy = TaxonomyVersion(
            version="abbott.v2",
            terms=TERMS,
            digest=compute_taxonomy_digest("abbott.v2", TERMS),
        )
        configuration = WorkflowConfiguration(
            "abbott.v2", "prompt.v2", "routing.v2", "b" * 40
        )
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = write_sources(Path(temporary))
            store = StatefulWorkflowStore(context(taxonomy=taxonomy))
            run = CanonicalWeeklyProposalService(
                store, configuration
            ).reconcile(registry1, registry2)
            receipt = CanonicalWeeklyProposalService(
                store, configuration
            ).classify(run.run_id, execute_llm=False)

        self.assertEqual(receipt.batch_id, 71)
        self.assertEqual(
            store.load_reconciliation_run(run.run_id).configuration.taxonomy_version,
            "abbott.v2",
        )


if __name__ == "__main__":
    unittest.main()
