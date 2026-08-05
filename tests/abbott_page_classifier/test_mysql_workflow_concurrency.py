"""Real two-connection MySQL concurrency contract for the Abbott workflow.

The test is opt-in because it requires a fresh ephemeral MySQL 8 instance with
only migrations 033, 045, 046, and 047 applied.  CI or a reviewer supplies only
the dedicated ``ABBOTT_TEST_MYSQL_*`` variables plus the exact explicit
ephemeral marker below; no populated or live database is accepted.
"""

from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
import hashlib
import json
import os
from pathlib import Path
import tempfile
import threading
import unittest

import mysql.connector
from openpyxl import Workbook

from agents.abbott_page_classifier.workflow_repository import MySqlWorkflowStore
from agents.abbott_page_classifier.workflow_service import (
    CanonicalWeeklyProposalService,
    WorkflowConfiguration,
)


EPHEMERAL_MARKER = "abbott-content-workflow-mysql-ephemeral-v1"
SEED_PREFIX = "codex-abbott-content-concurrency-v1"
BASELINE_MATERIAL_ID = f"{SEED_PREFIX}-baseline"
NEW_MATERIAL_ID = f"{SEED_PREFIX}-new"
BASELINE_URL = "https://abbottpro.ru/test/mysql-concurrency-baseline"
NEW_URL = "https://abbottpro.ru/test/mysql-concurrency-new"
SEED_RELEASE_KEY = f"{SEED_PREFIX}-active-release"
PROMPT_VERSION = "prompt.mysql-concurrency.v1"
ROUTING_VERSION = "routing.mysql-concurrency.v1"
CODE_REVISION = "c" * 40


def _configuration() -> WorkflowConfiguration:
    return WorkflowConfiguration(
        taxonomy_version="abbott.v1",
        prompt_version=PROMPT_VERSION,
        model_routing_version=ROUTING_VERSION,
        code_revision=CODE_REVISION,
    )


class MySqlWorkflowConcurrencyTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        required = {
            "host": os.environ.get("ABBOTT_TEST_MYSQL_HOST", "").strip(),
            "port": os.environ.get("ABBOTT_TEST_MYSQL_PORT", "").strip(),
            "database": os.environ.get("ABBOTT_TEST_MYSQL_DATABASE", "").strip(),
            "user": os.environ.get("ABBOTT_TEST_MYSQL_USER", "").strip(),
            "password": os.environ.get("ABBOTT_TEST_MYSQL_PASSWORD", ""),
        }
        if any(not value for value in required.values()):
            raise unittest.SkipTest("ephemeral MySQL variables are not configured")
        if required["database"] != "report_bd":
            raise unittest.SkipTest("concurrency test accepts only report_bd")
        if os.environ.get("ABBOTT_TEST_MYSQL_EPHEMERAL") != EPHEMERAL_MARKER:
            raise RuntimeError("explicit ephemeral MySQL marker is required")
        cls.settings = {
            **required,
            "port": int(required["port"]),
            "charset": "utf8mb4",
            "collation": "utf8mb4_unicode_ci",
        }

        connection = mysql.connector.connect(**cls.settings)
        cursor = connection.cursor()
        try:
            cursor.execute("SELECT DATABASE(), VERSION()")
            database, version = cursor.fetchone()
            if database != "report_bd" or not str(version).startswith("8."):
                raise RuntimeError("ephemeral MySQL 8 report_bd is required")
        finally:
            cursor.close()
            connection.close()

    def setUp(self):
        self.seed_snapshot_ids: tuple[int, int] = ()
        self.seed_release_id: int | None = None
        self._seed_isolated_predecessor()

    def tearDown(self):
        self._cleanup_test_owned_state()

    def _connect(self):
        return mysql.connector.connect(**self.settings)

    @staticmethod
    def _sha256(value: str) -> str:
        return hashlib.sha256(value.encode("utf-8")).hexdigest()

    @staticmethod
    def _all_base_tables(cursor) -> tuple[str, ...]:
        cursor.execute(
            "SELECT TABLE_NAME FROM information_schema.TABLES "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_TYPE = 'BASE TABLE' "
            "ORDER BY TABLE_NAME"
        )
        return tuple(str(row[0]) for row in cursor.fetchall())

    def _assert_fresh_migrations_only_database(self, cursor) -> None:
        tables = self._all_base_tables(cursor)
        required = {
            "portal_active_data_releases",
            "portal_content_approval_batches",
            "portal_content_catalog",
            "portal_content_reconciliation_runs",
            "portal_content_registry_entities",
            "portal_content_taxonomy_terms",
            "portal_content_taxonomy_versions",
            "portal_data_releases",
            "portal_dataset_snapshots",
        }
        if not required.issubset(tables):
            raise AssertionError("required Abbott migrations are not applied")
        for table in tables:
            if table in {
                "portal_content_taxonomy_terms",
                "portal_content_taxonomy_versions",
            }:
                continue
            cursor.execute(f"SELECT COUNT(*) FROM `{table}`")
            if int(cursor.fetchone()[0]) != 0:
                raise AssertionError(
                    f"ephemeral database is not fresh: {table} contains rows"
                )
        cursor.execute(
            "SELECT dataset_key, version, taxonomy_status "
            "FROM portal_content_taxonomy_versions ORDER BY id"
        )
        if cursor.fetchall() != [("abbott", "abbott.v1", "active")]:
            raise AssertionError("canonical taxonomy seed is not isolated")
        cursor.execute("SELECT COUNT(*) FROM portal_content_taxonomy_terms")
        if int(cursor.fetchone()[0]) != 40:
            raise AssertionError("canonical taxonomy seed is incomplete")

    def _seed_isolated_predecessor(self) -> None:
        connection = self._connect()
        cursor = connection.cursor()
        try:
            connection.start_transaction()
            self._assert_fresh_migrations_only_database(cursor)
            snapshot_ids = []
            for ordinal in (1, 2):
                digest = self._sha256(f"{SEED_PREFIX}:snapshot:{ordinal}")
                locator = f"test-only://{SEED_PREFIX}/snapshot/{ordinal}"
                cursor.execute(
                    """
                    INSERT INTO portal_dataset_snapshots (
                      snapshot_key, dataset_key, source_kind, source_locator,
                      content_sha256, content_bytes, source_row_count,
                      parser_version, import_status, imported_row_count,
                      rejected_row_count, private_archive_locator,
                      manifest_json, imported_at
                    ) VALUES (%s, 'abbott', %s, %s, %s, 0, 1,
                              %s, 'imported', 1, 0, %s, %s, NOW(6))
                    """,
                    (
                        f"{SEED_PREFIX}-snapshot-{ordinal}",
                        f"abbott_test_predecessor_{ordinal}",
                        locator,
                        digest,
                        f"{SEED_PREFIX}.parser.v1",
                        locator,
                        json.dumps(
                            {"owner": EPHEMERAL_MARKER, "ordinal": ordinal},
                            sort_keys=True,
                        ),
                    ),
                )
                snapshot_ids.append(int(cursor.lastrowid))
            cursor.execute(
                """
                INSERT INTO portal_data_releases (
                  dataset_key, release_key, source_snapshot_ids,
                  canonical_version_id, code_revision, release_status,
                  activated_at, activated_by
                ) VALUES ('abbott', %s, %s, %s, %s, 'active', NOW(), %s)
                """,
                (
                    SEED_RELEASE_KEY,
                    json.dumps(snapshot_ids),
                    f"{SEED_PREFIX}.canonical.v1",
                    CODE_REVISION,
                    EPHEMERAL_MARKER,
                ),
            )
            release_id = int(cursor.lastrowid)
            cursor.execute(
                """
                INSERT INTO portal_active_data_releases (
                  dataset_key, canonical_release_id, previous_release_id,
                  switched_at, switched_by, switch_reason
                ) VALUES ('abbott', %s, NULL, NOW(), %s, %s)
                """,
                (release_id, EPHEMERAL_MARKER, "ephemeral concurrency test"),
            )
            cursor.execute(
                """
                INSERT INTO portal_content_catalog (
                  canonical_release_id, source_snapshot_id, normalized_url,
                  normalized_url_hash, normalized_path, page_title,
                  material_id, material_type, source_slug, source_slug_hash,
                  access_label, is_active, source_sheet,
                  source_row_ordinal, source_row_fingerprint, section_key,
                  direction_key, valid_from
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, 'articles',
                          %s, %s, 'all', 1, %s, 1, %s, %s,
                          'cardiology', '2026-01-01 00:00:00')
                """,
                (
                    release_id,
                    snapshot_ids[0],
                    BASELINE_URL,
                    self._sha256(BASELINE_URL),
                    "/test/mysql-concurrency-baseline",
                    "Ephemeral MySQL concurrency baseline",
                    BASELINE_MATERIAL_ID,
                    "mysql-concurrency-baseline",
                    self._sha256("mysql-concurrency-baseline"),
                    SEED_PREFIX,
                    self._sha256(f"{SEED_PREFIX}:catalog:1"),
                    SEED_PREFIX,
                ),
            )
            connection.commit()
            self.seed_snapshot_ids = (snapshot_ids[0], snapshot_ids[1])
            self.seed_release_id = release_id
        except Exception:
            connection.rollback()
            raise
        finally:
            cursor.close()
            connection.close()

    @staticmethod
    def _delete_ids(cursor, table: str, column: str, values: set[int]) -> None:
        if not values:
            return
        placeholders = ", ".join(["%s"] * len(values))
        cursor.execute(
            f"DELETE FROM `{table}` WHERE `{column}` IN ({placeholders})",
            tuple(sorted(values)),
        )

    def _cleanup_test_owned_state(self) -> None:
        if self.seed_release_id is None:
            return
        connection = self._connect()
        cursor = connection.cursor()
        try:
            connection.start_transaction()
            cursor.execute(
                """
                SELECT id, registry1_snapshot_id, registry2_snapshot_id,
                       predecessor_release_id, prompt_version,
                       model_routing_version, code_revision
                FROM portal_content_reconciliation_runs
                ORDER BY id FOR UPDATE
                """
            )
            run_rows = tuple(cursor.fetchall())
            if any(
                int(row[3]) != self.seed_release_id
                or str(row[4]) != PROMPT_VERSION
                or str(row[5]) != ROUTING_VERSION
                or str(row[6]) != CODE_REVISION
                for row in run_rows
            ):
                raise AssertionError("non-test reconciliation state detected")
            run_ids = {int(row[0]) for row in run_rows}
            run_snapshot_ids = {
                int(snapshot_id)
                for row in run_rows
                for snapshot_id in row[1:3]
            }
            cursor.execute(
                "SELECT id, reconciliation_run_id FROM portal_content_approval_batches "
                "ORDER BY id FOR UPDATE"
            )
            batch_rows = tuple(cursor.fetchall())
            if any(int(row[1]) not in run_ids for row in batch_rows):
                raise AssertionError("non-test approval batch detected")
            batch_ids = {int(row[0]) for row in batch_rows}
            cursor.execute(
                "SELECT id, material_id FROM portal_content_registry_entities "
                "ORDER BY id FOR UPDATE"
            )
            entity_rows = tuple(cursor.fetchall())
            if any(
                str(row[1]) not in {BASELINE_MATERIAL_ID, NEW_MATERIAL_ID}
                for row in entity_rows
            ):
                raise AssertionError("non-test registry entity detected")
            entity_ids = {int(row[0]) for row in entity_rows}
            cursor.execute(
                "SELECT id, content_entity_id FROM portal_content_registry_aliases "
                "ORDER BY id FOR UPDATE"
            )
            if any(int(row[1]) not in entity_ids for row in cursor.fetchall()):
                raise AssertionError("non-test registry alias detected")
            cursor.execute(
                "SELECT id, content_entity_id FROM portal_content_classification_events "
                "ORDER BY id FOR UPDATE"
            )
            if any(int(row[1]) not in entity_ids for row in cursor.fetchall()):
                raise AssertionError("non-test classification event detected")
            cursor.execute(
                "SELECT id, reconciliation_run_id FROM portal_content_reconciliation_items "
                "ORDER BY id FOR UPDATE"
            )
            if any(int(row[1]) not in run_ids for row in cursor.fetchall()):
                raise AssertionError("non-test reconciliation item detected")
            cursor.execute(
                "SELECT id, reconciliation_run_id FROM portal_content_llm_attempts "
                "ORDER BY id FOR UPDATE"
            )
            if any(int(row[1]) not in run_ids for row in cursor.fetchall()):
                raise AssertionError("non-test LLM attempt detected")
            cursor.execute(
                "SELECT id, approval_batch_id FROM portal_content_approval_items "
                "ORDER BY id FOR UPDATE"
            )
            if any(int(row[1]) not in batch_ids for row in cursor.fetchall()):
                raise AssertionError("non-test approval item detected")
            cursor.execute(
                "SELECT id, canonical_release_id FROM portal_content_catalog "
                "ORDER BY id FOR UPDATE"
            )
            if any(int(row[1]) != self.seed_release_id for row in cursor.fetchall()):
                raise AssertionError("non-test catalog row detected")
            cursor.execute(
                "SELECT id, release_key FROM portal_data_releases ORDER BY id FOR UPDATE"
            )
            if cursor.fetchall() != [(self.seed_release_id, SEED_RELEASE_KEY)]:
                raise AssertionError("non-test release detected")
            cursor.execute(
                "SELECT canonical_release_id FROM portal_active_data_releases "
                "ORDER BY dataset_key FOR UPDATE"
            )
            if cursor.fetchall() != [(self.seed_release_id,)]:
                raise AssertionError("non-test active pointer detected")
            expected_snapshot_ids = set(self.seed_snapshot_ids) | run_snapshot_ids
            cursor.execute(
                "SELECT id FROM portal_dataset_snapshots ORDER BY id FOR UPDATE"
            )
            if {int(row[0]) for row in cursor.fetchall()} != expected_snapshot_ids:
                raise AssertionError("non-test source snapshot detected")

            self._delete_ids(
                cursor, "portal_content_classification_events",
                "content_entity_id", entity_ids,
            )
            self._delete_ids(
                cursor, "portal_content_approval_items", "approval_batch_id", batch_ids
            )
            self._delete_ids(
                cursor, "portal_content_approval_batches", "id", batch_ids
            )
            self._delete_ids(
                cursor, "portal_content_llm_attempts", "reconciliation_run_id", run_ids
            )
            self._delete_ids(
                cursor, "portal_content_reconciliation_items",
                "reconciliation_run_id", run_ids,
            )
            self._delete_ids(
                cursor, "portal_content_reconciliation_runs", "id", run_ids
            )
            self._delete_ids(
                cursor, "portal_content_registry_aliases", "content_entity_id", entity_ids
            )
            self._delete_ids(
                cursor, "portal_content_registry_entities", "id", entity_ids
            )
            cursor.execute(
                "DELETE FROM portal_content_catalog WHERE canonical_release_id = %s",
                (self.seed_release_id,),
            )
            cursor.execute(
                "DELETE FROM portal_active_data_releases "
                "WHERE dataset_key = 'abbott' AND canonical_release_id = %s",
                (self.seed_release_id,),
            )
            cursor.execute(
                "DELETE FROM portal_data_releases "
                "WHERE id = %s AND release_key = %s",
                (self.seed_release_id, SEED_RELEASE_KEY),
            )
            self._delete_ids(
                cursor, "portal_dataset_snapshots", "id", expected_snapshot_ids
            )
            self._assert_fresh_migrations_only_database(cursor)
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        finally:
            cursor.close()
            connection.close()

    @staticmethod
    def _sources(directory: Path) -> tuple[Path, Path]:
        registry1 = directory / "registry1.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "кардио"
        sheet.append(("ID", "Название", "ссылка", "Направление", "Тип контента"))
        sheet.append((
            NEW_MATERIAL_ID,
            "Concurrent material",
            NEW_URL,
            "Кардиология",
            "Статьи",
        ))
        workbook.save(registry1)
        registry2 = directory / "registry2.csv"
        registry2.write_text(
            "ID,Название,URL,Направление,Тип материала\n"
            f"{NEW_MATERIAL_ID},Concurrent material,{NEW_URL},"
            "Кардиология,Статьи\n",
            encoding="utf-8",
        )
        return registry1, registry2

    def test_two_connections_reconcile_classify_create_and_finalize_once(self):
        # Exercise the actual connector's ``%s`` binding before workflow SQL.
        connector = self._connect()
        cursor = connector.cursor()
        cursor.execute("SELECT %s + %s", (20, 22))
        self.assertEqual(cursor.fetchone()[0], 42)
        cursor.close()
        connector.close()

        reconcile_barrier = threading.Barrier(2)
        classify_barrier = threading.Barrier(2)
        with tempfile.TemporaryDirectory() as temporary:
            registry1, registry2 = self._sources(Path(temporary))

            def worker():
                service = CanonicalWeeklyProposalService(
                    MySqlWorkflowStore(self._connect),
                    _configuration(),
                )
                reconcile_barrier.wait(timeout=15)
                reconciled = service.reconcile(registry1, registry2)
                classify_barrier.wait(timeout=15)
                classified = service.classify(
                    reconciled.run_id, execute_llm=False
                )
                return reconciled, classified

            with ThreadPoolExecutor(max_workers=2) as pool:
                futures = [pool.submit(worker) for _ in range(2)]
                results = [future.result(timeout=45) for future in futures]

        run_ids = {value[0].run_id for value in results}
        batch_ids = {value[1].batch_id for value in results}
        self.assertEqual(len(run_ids), 1)
        self.assertEqual(len(batch_ids), 1)
        run_id = next(iter(run_ids))
        batch_id = next(iter(batch_ids))
        self.assertGreater(run_id, 0)
        self.assertGreater(batch_id, 0)

        connection = self._connect()
        cursor = connection.cursor()
        cursor.execute(
            "SELECT run_status, COUNT(*) FROM portal_content_reconciliation_runs "
            "WHERE id = %s GROUP BY run_status",
            (run_id,),
        )
        self.assertEqual(cursor.fetchall(), [("finalized", 1)])
        cursor.execute(
            "SELECT COUNT(*), MIN(id), MAX(id) "
            "FROM portal_content_registry_entities WHERE material_id = %s",
            (NEW_MATERIAL_ID,),
        )
        entity_count, minimum_id, maximum_id = cursor.fetchone()
        self.assertEqual(entity_count, 1)
        self.assertEqual(minimum_id, maximum_id)
        cursor.execute(
            "SELECT COUNT(*) FROM portal_content_registry_aliases "
            "WHERE alias_type = 'material_id' AND alias_value = %s",
            (NEW_MATERIAL_ID,),
        )
        self.assertEqual(cursor.fetchone()[0], 1)
        cursor.execute(
            "SELECT COUNT(*) FROM portal_content_approval_batches "
            "WHERE reconciliation_run_id = %s",
            (run_id,),
        )
        self.assertEqual(cursor.fetchone()[0], 1)
        cursor.execute(
            "SELECT COUNT(*) FROM portal_content_approval_batches WHERE id = %s",
            (batch_id,),
        )
        self.assertEqual(cursor.fetchone()[0], 1)
        cursor.close()
        connection.close()


if __name__ == "__main__":
    unittest.main()
